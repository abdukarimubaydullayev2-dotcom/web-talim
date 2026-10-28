"""Shared Word formatting for learning modules. No build side effects."""
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH

def plain_heading(s):return re.sub(r'[^\w\s‘’]', ' ',s).strip()

def new_doc(module=7):
 d=Document();sec=d.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11);sec.top_margin=sec.bottom_margin=Inches(.7);sec.left_margin=sec.right_margin=Inches(.75)
 for style in d.styles:
  if style.type in [1,2]:
   style.font.name='Calibri';style.font.size=Pt(14);style.font.color.rgb=RGBColor(0,0,0)
  if style.type==1:
   p=style.paragraph_format;p.line_spacing=1;p.space_after=Pt(4)
   if style.element.pPr is not None:
    for border in list(style.element.pPr.findall(qn('w:pBdr'))):style.element.pPr.remove(border)
 for name in ['Title','Heading 1','Heading 2','Heading 3']:
  d.styles[name].font.bold=True;d.styles[name].paragraph_format.keep_with_next=True
  d.styles[name].paragraph_format.space_before=Pt(6)
  d.styles[name].paragraph_format.space_after=Pt(3)
 d.core_properties.author='Ubaydullayev Abdukarim'
 foot=d.sections[0].footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
 foot.add_run(f'Web Ta’lim  |  {module} modul  |  ')
 field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');foot._p.append(field)
 return d

def table_style(table,widths):
 table.autofit=False
 for column,width in zip(table.columns,widths):column.width=Inches(width)
 for row in table.rows:
  for i,cell in enumerate(row.cells):
   cell.width=Inches(widths[i]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   pr=cell._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');b.append(e)
   pr.append(b);m=OxmlElement('w:tcMar')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:w'),'90');e.set(qn('w:type'),'dxa');m.append(e)
   pr.append(m)
   if row is table.rows[0]:pass
   for p in cell.paragraphs:
    p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=1
    if i in [0,len(widths)-1]:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
  row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
 header=table.rows[0]
 for cell in header.cells:
  sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E8EEF3');cell._tc.get_or_add_tcPr().append(sh)
  for run in cell.paragraphs[0].runs:run.bold=True
 header._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
