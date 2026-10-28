"""Build the 180-hour Backend module, Word work plan and lesson notes.
Run with the primary document runtime Python. Content lives in backend_units.py.
"""
import copy
import html
import json
import re
import zipfile
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from backend_units import UNITS
from backend_cases import CASES

ROOT=Path(__file__).resolve().parents[1]
DIST=ROOT/'dist'
MAT=DIST/'materials'
MAT.mkdir(exist_ok=True)
E=lambda x:html.escape(str(x),quote=True)
GROUP_NAMES=['Backend va HTTP asoslari','Node js modullar va asinxron ishlash','Express va REST API','SQL SQLite va doimiy saqlash','Autentifikatsiya sessiya va ruxsatlar','Validatsiya xavfsizlik va yaxlitlik','Frontend bilan bog‘lash va testlash','Joylashtirish va xizmatni boshqarish','Yakuniy Backend loyiha']
CONTEXTS=[
'Kutubxona misolida brauzer kitob nomini ko‘rsatadi, server esa qaysi yozuvni qaytarishni va qanday qoidaga amal qilishni belgilaydi. So‘rovni jo‘natish bilan uni bajarishga ruxsat berish farq qiladi. Bir foydalanuvchining noto‘g‘ri yuborgan qiymati boshqa foydalanuvchi ma’lumotini buzmasligi kerak.',
'Node muhitida frontenddan tanish JavaScript sintaksisi saqlanadi, biroq bajarilish muhiti boshqa. Fayl tizimi, jarayon va tarmoq server tomonida ishlaydi. Mashqni alohida katalogda bajarib, kod qaysi fayl va qaysi jarayonga ta’sir qilayotganini kuzatish kerak.',
'HTTP API yaratishda har route faqat muvaffaqiyatli holatni emas, yaroqsiz kirish va topilmagan resursni ham tushunarli ifodalashi kerak. So‘rovning bosqichlari ketma-ketlikda ishlaydi: kirishni o‘qish, tekshirish, kerakli amal va javob. Har bosqichning chegarasi aniq bo‘lsa, xato sababini topish osonlashadi.',
'Vazifalar ro‘yxati faqat server xotirasida qolsa, jarayon to‘xtaganda yo‘qoladi. Relatsion baza yozuvni saqlash bilan birga ustunlar va bog‘lanishlarga qoida beradi. API tekshiruvi va baza cheklovlari birgalikda ishlaydi. SQL misollarini haqiqiy ma’lumotdan ajratilgan demo bazada bajaring.',
'Kirish tizimida kimlikni tekshirish va amal uchun ruxsat berish alohida qarorlardir. Foydalanuvchi paroli himoyalangan hosila bilan tekshiriladi, keyingi so‘rovlarda server sessiyasi ishlatiladi. Brauzer ko‘rsatgan tugmaga emas, serverdagi tekshiruvga tayaniladi. Faqat sun’iy mashq hisoblaridan foydalaning.',
'API chegarasiga kelgan qiymatni ishonchli deb olish mumkin emas. Bir tekshiruv barcha muammolarni hal qilmaydi: format, biznes qoida, SQL qiymati va brauzerda chiqarish har biri o‘z joyida boshqariladi. Natija rad etilganda mijoz nima tuzatishini tushunishi, server esa ichki tafsilotlarni oshkor qilmasligi kerak.',
'Frontend bilan server o‘rtasidagi kelishuv ishlaganda ham xato va bo‘sh holatlar uchraydi. Testlar odatiy natija bilan cheklanmay, cheklov va ruxsatni ham tekshiradi. Alohida funksiya sinovi bilan HTTP orqali tizim qismlarining birga ishlash sinovi bir-birini to‘ldiradi.',
'Lokal kompyuterda ishlagan ilova hostingda ham xuddi shunday muhit va disk sharoitiga ega bo‘lavermaydi. Runtime, port, HTTPS, origin, sirlar va doimiy ombor oldindan tekshiriladi. Xizmat ishlashini kuzatish, zaxira olish va alohida joyda tiklash ham dasturning ish jarayoniga kiradi.',
'Yakuniy loyiha oldin o‘tilgan qismlarni bir tizimga birlashtiradi. Maqsad tayyor kodni ko‘chirish bilangina cheklanish emas: talabning qaysi route, tekshiruv va test bilan bajarilishini tushuntirish kerak. Ishlaydigan kichik loyiha, qabul mezonlari va ochiq yozilgan cheklovlar himoya uchun asos bo‘ladi.'
]
SOURCES={
1:[('Node js asoslari','https://nodejs.org/en/learn/getting-started/introduction-to-nodejs'),('HTTP ma’lumotnomasi','https://developer.mozilla.org/en-US/docs/Web/HTTP')],
2:[('Node js o‘quv bo‘limi','https://nodejs.org/en/learn/getting-started/introduction-to-nodejs'),('Node fayl tizimi','https://nodejs.org/api/fs.html')],
3:[('Express routing','https://expressjs.com/en/guide/routing.html'),('Express middleware','https://expressjs.com/en/guide/using-middleware/'),('Express xatolari','https://expressjs.com/en/guide/error-handling.html')],
4:[('SQLite SQL tili','https://www.sqlite.org/lang.html'),('Node SQLite','https://nodejs.org/api/sqlite.html')],
5:[('OWASP autentifikatsiya','https://cheatsheetseries.owasp.org/cheatsheets/Authentication_Cheat_Sheet.html'),('OWASP parol saqlash','https://cheatsheetseries.owasp.org/cheatsheets/Password_Storage_Cheat_Sheet.html'),('Express sessiya','https://expressjs.com/en/resources/middleware/session.html')],
6:[('OWASP kirish tekshiruvi','https://cheatsheetseries.owasp.org/cheatsheets/Input_Validation_Cheat_Sheet.html'),('SQLite tranzaksiya','https://www.sqlite.org/lang_transaction.html')],
7:[('Node test','https://nodejs.org/api/test.html'),('MDN fetch','https://developer.mozilla.org/en-US/docs/Web/API/Fetch_API/Using_Fetch')],
8:[('Node LTS relizlari','https://nodejs.org/en/about/previous-releases'),('Express ishlab chiqarish xavfsizligi','https://expressjs.com/en/advanced/best-practice-security.html'),('SQLite backup','https://www.sqlite.org/backup.html')],
9:[('Express qo‘llanmasi','https://expressjs.com/en/guide/routing.html'),('Node SQLite','https://nodejs.org/api/sqlite.html'),('OWASP sessiya','https://cheatsheetseries.owasp.org/cheatsheets/Session_Management_Cheat_Sheet.html')]
}
MODE_HELP={
'report':'Bu sxema yoki tahlil namunasi. Uni kod sifatida bajarish talab etilmaydi; yozma hisobot va kuzatuv tayyorlanadi.',
'terminal':'Buyruqlar lokal loyiha katalogidagi terminalda bajariladi. Kommentariylarni alohida o‘qing; buyruq va fayl mazmunini chalkashtirmang.',
'server':'Bu Node server namunasi. Faylni lokal terminalda ishga tushiring, keyin ko‘rsatilgan localhost manzilini brauzerda oching. Konsol bajaruvchisi serverning tashqi HTTP manzilini ochib berishi kafolatlanmaydi.',
'files':'Namuna izohlarida ko‘rsatilgan alohida fayllarga bo‘linadi. So‘ng asosiy fayl terminaldan ishga tushiriladi.',
'node':'Bu mustaqil Node konsol misoli. Node 24 LTSda lokal ishga tushirish asosiy tekshiruv usuli; onlayn xizmatdagi Node versiyasi va modullar farq qilishi mumkin.',
'snippet':'Bu oldingi loyiha ichiga qo‘shiladigan kod parchasi. app, db, server yoki jadvallar oldindan tayyorlangan bo‘lishi kerak; parcha yolg‘iz fayl sifatida to‘liq ilova emas.',
'sql':'SQLni alohida demo bazada bajaring. Kerakli jadval va boshlang‘ich yozuvlar oldindan yaratiladi; Nodega ulash keyingi mavzularda ko‘rsatiladi.',
'browser':'Bu brauzerda, ishlayotgan lokal Node loyihasining frontendida bajariladi. Uni Node konsolida yolg‘iz bajarish DOM yoki nisbiy URL xatosiga olib kelishi mumkin.'}

VIDEOS={}
def video(units,ident,title,channel,coverage,evidence=None):
 for u in units:
  VIDEOS.setdefault(u,[]).append(dict(id=ident,title=title,channel=channel,language='O‘zbekcha',coverage=coverage,task='Videodagi asosiy tushunchani konspekt bilan solishtiring. Hozirgi darsda o‘tilgan qismni tushuntiring yoki amaliy topshiriqda qo‘llang.',source='https://www.youtube.com/watch?v='+ident,verifiedFrom=evidence or 'https://www.youtube.com/watch?v='+ident,checkedOn='2026-10-27'))
video([4],'69pxFJF5SkY','NodeJS — Kirish','Farkhod Dadajanov','Node.jsga kirish uchun aloqador dars. Video eski relizga tegishli bo‘lishi mumkin; hozirgi kurs Node 24 LTS va konspektdagi buyruqlarga tayanadi.')
video([7],'PiYJ98YZqM0','Node.js asoslari — FS moduli','Bexruz Xolmuminov','Fayl tizimi haqida dars. Videodagi callback yoki modul usuli farq qilsa, konspektdagi fs/promises va CommonJS misolini asos qiling.')
video([11],'CUlD5c0FKA0','Express.js — Asosiy tushunchalar','Bexruz Xolmuminov','Expressga kirish darsi. Video eski versiyaga oid; kurs misollari Express 5 asosida tuzilgan.')
video([11,13],'3pnIAhr9lS8','Express.js — JSON response va middlewarelar','Bexruz Xolmuminov','JSON javobi va middleware tushuntiriladi. Bu barcha validatsiya qoidalarining to‘liq darsi emas; ularni konspektdan davom ettiring.')
video([17],'2L4dye3x1Tk','SQL — DB Browser for SQLite','Sanjar Suvonov','SQLite va DB Browser bilan tanishishga aloqador dars. Muallifning Android kursidan; ushbu mavzuda faqat SQLite vositasi qismi kerak, Android topshirig‘i talab etilmaydi.','https://t.me/s/mobileracademy?after=117')
video([18],'oekR8yp4CtE','SQL — SELECT buyrug‘i, 1-qism','ITOne Uz','Umumiy SQL SELECT tushunchasi uchun. Videoda MySQL muhiti ishlatiladi; SQLite sintaksisi uchun konspektni asos qiling.','https://t.me/s/itone_uz?before=316')
video([18],'_KJzMkffd_k','SQL — SELECT va WHERE, 2-qism','ITOne Uz','SQLda WHERE bilan tanlash. MySQL va SQLite farqlarida konspektdagi SQLite misoli ustun.','https://t.me/s/itone_uz?before=316')
video([20],'FX4AFqP9Jow','SQL — JOIN va jadvallarni bog‘lash','ITOne Uz','JOIN tushunchasiga aloqador dars. Video MySQL muhitida; SQLite foreign key va COUNT qoidalari konspektda alohida berilgan.','https://t.me/s/itone_uz?before=316')

def quiz_for(u,index,n):
 qs=[]
 for k,(term,definition) in enumerate(u['concepts']):
  distract=[x[1] for j,x in enumerate(u['concepts']) if j!=k][:3]
  qs.append(dict(q=f'{term} bo‘yicha qaysi izoh to‘g‘ri?',options=[definition]+distract,correct=0,explanation=definition))
 for raw in CASES[index]:
  q,*opts=raw.split('~');assert len(opts)==4
  qs.append(dict(q=q,options=opts,correct=0,explanation=opts[0]+'.'))
 others=[UNITS[(index+j)%len(UNITS)] for j in [1,7,13]]
 qs.append(dict(q='Konspektdagi namuna uchun qaysi natija kutiladi?',options=[u['expected']]+[x['expected'] for x in others],correct=0,explanation=u['expected']))
 qs.append(dict(q='Ushbu darsdagi namuna uchun qaysi xatoga alohida e’tibor beriladi?',options=[u['pitfall']]+[x['pitfall'] for x in others],correct=0,explanation=u['pitfall']))
 # Deliberate distinction between environments is central to Backend practice.
 help_correct=MODE_HELP[u['mode']]
 pool=[v for v in MODE_HELP.values() if v!=help_correct][:3]
 qs.append(dict(q='Ushbu darsning namunasini qaysi muhitda va qanday tartibda qo‘llash kerak?',options=[help_correct]+pool,correct=0,explanation=help_correct))
 assert len(qs)==10
 for i,q in enumerate(qs):
  shift=(n+i)%4;q['options']=q['options'][shift:]+q['options'][:shift];q['correct']=(-shift)%4
 return qs

def plain_heading(s):return re.sub(r'[^\w\s‘’]', ' ',s).strip()
def new_doc():
 d=Document();sec=d.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11);sec.top_margin=sec.bottom_margin=Inches(.7);sec.left_margin=sec.right_margin=Inches(.75)
 for style in d.styles:
  if style.type in [1,2]:
   style.font.name='Calibri';style.font.size=Pt(14);style.font.color.rgb=RGBColor(0,0,0)
  if style.type==1:
   p=style.paragraph_format;p.line_spacing=1;p.space_after=Pt(4)
   if style.element.pPr is not None:
    for border in list(style.element.pPr.findall(qn('w:pBdr'))):style.element.pPr.remove(border)
 for name in ['Title','Heading 1','Heading 2','Heading 3']:
  d.styles[name].font.bold=True;d.styles[name].paragraph_format.keep_with_next=True
  d.styles[name].paragraph_format.space_before=Pt(6)
  d.styles[name].paragraph_format.space_after=Pt(3)
 d.core_properties.author='Ubaydullayev Abdukarim'
 foot=d.sections[0].footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
 foot.add_run('Web Ta’lim  |  6 modul  |  ')
 field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');foot._p.append(field)
 return d

def make_note(d,r,u,practical):
 # Two longer practice notes need tighter paragraph spacing to keep their
 # bibliography with the reading material. Font size and line spacing stay 14/1.
 if r['number'] in [48,62]:d.styles['Normal'].paragraph_format.space_after=Pt(2)
 d.add_paragraph(plain_heading(r['title']),style='Title')
 d.add_paragraph(f'6 modul  |  {r["number"]} mavzu  |  2 soat  |  {r["kind"]}')
 d.add_paragraph('Ubaydullayev Abdukarim  |  Namangan shahar 3 sonli texnikum\n3 IT  |  2026/2027')
 d.add_paragraph('Dars maqsadi',style='Heading 1');d.add_paragraph(r['outcome'])
 d.add_paragraph('Tayanch bilimlar',style='Heading 2');d.add_paragraph(r['prerequisite'])
 d.add_paragraph('Mavzuning mazmuni',style='Heading 1');d.add_paragraph('Oldingi nazariy darsdagi namunaga tayangan holda quyidagi vazifani bajaring. Ishning bajarilish muhiti, kutiladigan natija va xato holatini dalil bilan yozing.' if practical else CONTEXTS[u['group']-1])
 for term,definition in u['concepts']:
  d.add_paragraph(plain_heading(term),style='Heading 2');d.add_paragraph(definition)
 d.add_paragraph('Namuna va ishlash muhiti',style='Heading 1');d.add_paragraph(MODE_HELP[u['mode']])
 for line in u['code'].splitlines():
  p=d.add_paragraph(line or ' ');p.paragraph_format.space_after=Pt(0)
 d.add_paragraph('Namuna natijasi',style='Heading 2');d.add_paragraph(u['expected'])
 d.add_paragraph('Ko‘p uchraydigan xato',style='Heading 2');d.add_paragraph(u['pitfall'])
 d.add_paragraph('Amaliy bajarish tartibi' if practical else 'Mavzuni mustahkamlash',style='Heading 1')
 d.add_paragraph(u['task'] if practical else 'Asosiy tushunchalarni izohlang, namunadagi kirish va natijani ajrating. Kodni kuzatish bu darsning qismi; mustaqil amaliy ish keyingi darsda bajariladi.')
 if practical:
  for i,x in enumerate(r['assignment']['steps']):
   if i!=1:d.add_paragraph(x,style='List Number')
 d.add_paragraph('Tekshirish mezonlari',style='Heading 2')
 for x in r['criteria']:d.add_paragraph(x,style='List Bullet')
 d.add_paragraph('Atamalar',style='Heading 1')
 for term,definition in u['concepts']:d.add_paragraph(term+' — '+definition.split('. ')[0]+'.')
 d.add_paragraph('Manbalar',style='Heading 1')
 for title,url in SOURCES[u['group']]:
  p=d.add_paragraph();link=OxmlElement('w:hyperlink')
  rid=d.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
  link.set(qn('r:id'),rid);run=OxmlElement('w:r');pr=OxmlElement('w:rPr')
  font=OxmlElement('w:rFonts');font.set(qn('w:ascii'),'Calibri');font.set(qn('w:hAnsi'),'Calibri');pr.append(font)
  size=OxmlElement('w:sz');size.set(qn('w:val'),'28');pr.append(size)
  run.append(pr);text=OxmlElement('w:t');text.text=title;run.append(text);link.append(run);p._p.append(link)

def table_style(table,widths):
 table.autofit=False
 for column,width in zip(table.columns,widths):column.width=Inches(width)
 for row in table.rows:
  for i,cell in enumerate(row.cells):
   cell.width=Inches(widths[i]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   pr=cell._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');b.append(e)
   pr.append(b);m=OxmlElement('w:tcMar')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:w'),'90');e.set(qn('w:type'),'dxa');m.append(e)
   pr.append(m)
   if row is table.rows[0]:pass
   for p in cell.paragraphs:
    p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=1
    if i in [0,len(widths)-1]:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
  row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
 header=table.rows[0]
 for cell in header.cells:
  sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E8EEF3');cell._tc.get_or_add_tcPr().append(sh)
  for run in cell.paragraphs[0].runs:run.bold=True
 header._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))

assert len(UNITS)==len(CASES)==45
DATA={};GROUPS=[]
for group_no,name in enumerate(GROUP_NAMES,1):GROUPS.append(dict(number=group_no,name=name,hours=0,topics=[]))
for index,u in enumerate(UNITS):
 unit_no=index+1
 if unit_no==42:u['mode']='snippet'
 for practical in [False,True]:
  n=index*2+1+int(practical);title=('Amaliy mashg‘ulot — ' if practical else '')+u['title']
  terms=[x[0] for x in u['concepts']]
  outcome=(f'{terms[0]} va {terms[1]} tushunchalarini izohlaydi; {terms[2]} qo‘llanishini namunada ajratadi.' if not practical else u['task'])
  prerequisite=('5-moduldagi JavaScript, obyekt, funksiya, JSON, fetch, async/await va Git asoslari.' if n==1 else f'6-modulning 1–{n-1}-mavzulari. '+(f'Asosiy tayanch: {n-1}-mavzudagi tushuntirish.' if practical else 'Oldingi dars natijasini eslang; yangi tushunchalar quyida tushuntiriladi.'))
  criteria=[f'{terms[0]} va {terms[1]} vazifasi misol bilan tushuntirilgan.', 'Namuna muhiti va kutiladigan natija to‘g‘ri ajratilgan.','Ko‘rsatilgan xato holati tekshirilgan va sababi tushuntirilgan.','Natija yoki xato sababi dalil bilan yozilgan.']
  qa=[[f'{term} nimani anglatadi va bu darsda qanday vazifa bajaradi?',definition] for term,definition in u['concepts']]
  qa += [['Namuna natijasini qanday izohlaysiz?',u['expected']],['Ushbu misolda qaysi xatoni oldini olish kerak?',u['pitfall']],['Namunani qaysi muhitda ishlatish kerak?',MODE_HELP[u['mode']]],['Mustaqil ishning mazmuni nima?',u['task']],['Natija qanday dalil bilan tekshiriladi?','Namuna natijasi bilan olingan natija solishtiriladi. Tegishli darsda noto‘g‘ri qiymat yoki ruxsat holati ham sinaladi. Faqat kodni ko‘chirish bajarilgan ish dalili emas.']]
  assignment=None
  if practical:
   assignment=dict(mode=u['mode'],task=u['task'],code=u['code'],expected=u['expected'],draftId=f'm6-r30-{n}',unit=unit_no,
    steps=['Oldingi nazariy darsni va quyidagi namuna izohini o‘qing.',MODE_HELP[u['mode']],'Namuna natijasini tekshiring, keyin topshiriqdagi o‘zgarishni bajaring.','Kamida bitta oddiy va mavzuga tegishli bitta salbiy yoki chegaraviy holatni tekshiring.','Kuzatilgan natija, xato va tuzatishni hisobotga yozing.'],criteria=criteria)
  rawhtml='<h3>Dars maqsadi</h3><p>'+E(outcome)+'</p><h3>Tayanch bilimlar</h3><p>'+E(prerequisite)+'</p><h3>Mavzuning mazmuni</h3><p>'+E(CONTEXTS[u['group']-1])+'</p>'
  rawhtml+=''.join('<h3>'+E(t)+'</h3><p>'+E(defn)+'</p>' for t,defn in u['concepts'])
  rawhtml+='<h3>Namuna va ishlash muhiti</h3><p>'+E(MODE_HELP[u['mode']])+'</p><pre class="lesson-code"><code>'+E(u['code'])+'</code></pre><h3>Namuna natijasi</h3><p>'+E(u['expected'])+'</p><h3>Ko‘p uchraydigan xato</h3><p>'+E(u['pitfall'])+'</p>'
  rawhtml+='<h3>'+('Amaliy vazifa' if practical else 'Mustahkamlash')+'</h3><p>'+E(u['task'] if practical else 'Beshta tushunchani o‘z so‘zingiz bilan izohlang. Namunani kuzating; mustaqil amaliy vazifa keyingi darsda beriladi.')+'</p>'
  rawhtml+='<h3>Tekshirish mezonlari</h3><ul>'+''.join('<li>'+E(x)+'</li>' for x in criteria)+'</ul><h3>Atamalar</h3>'
  rawhtml+=''.join('<p><b>'+E(t)+'</b> — '+E(d.split('. ')[0])+'.</p>' for t,d in u['concepts'])
  rawhtml+='<h3>Manbalar</h3><ul>'+''.join(f'<li><a href="{url}" target="_blank" rel="noopener">{E(t)}</a></li>' for t,url in SOURCES[u['group']])+'</ul>'
  slides=[dict(title=E(title),body=[f'6-modul · {n}-mavzu · 2 soat',E('Amaliy mashg‘ulot' if practical else 'Nazariy mashg‘ulot')]),dict(title='Dars maqsadi',body=[E(outcome)]),dict(title='Asosiy tushunchalar',body=[E(' · '.join(terms))],kind='agenda')]
  for t,d in u['concepts']:slides.append(dict(title=E(t),body=[E(x+'.') for x in d.rstrip('.').split('. ')],kind='example'))
  slides += [dict(title='Namuna',body=[],code=E(u['code']),kind='example'),dict(title='Namuna natijasi',body=[E(u['expected']),E(MODE_HELP[u['mode']])],kind='solution'),dict(title='Xatoni aniqlang',body=[E(u['pitfall'])],kind='review'),dict(title='Bajariladigan ish' if practical else 'Mustahkamlash',body=[E(u['task'] if practical else 'Tushunchalarni misol bilan izohlang. Mustaqil kod yozish keyingi amaliy darsda bajariladi.')],kind='task'),dict(title='Atamalarni takrorlang',body=[E(t+' — '+d.split('. ')[0]+'.') for t,d in u['concepts']],kind='terms'),dict(title='Tekshirish',body=[E(x) for x in criteria],kind='review'),dict(title='Manbalar',body=[],links=[dict(title=E(t),url=url) for t,url in SOURCES[u['group']]])]
  quiz=quiz_for(u,index,n)
  # All renderer-facing strings are escaped; code and authoring source remain raw.
  for q in quiz:
   q['q']=E(q['q']);q['options']=[E(v) for v in q['options']];q['explanation']=E(q['explanation'])
  r=dict(number=n,title=title,kind='Amaliy' if practical else 'Nazariy',type='amaliy' if practical else 'nazariy',hours=2,unit=unit_no,
    outcome=outcome,prerequisite=prerequisite,prerequisites=[] if n==1 else [n-1],criteria=criteria,html=rawhtml,slides=slides,qa=[[E(q),E(a)] for q,a in qa],quiz=quiz,
    download=f'materials/m6-{n:03d}-konspekt.docx',videos=copy.deepcopy(VIDEOS.get(unit_no,[])))
  if assignment:r['assignment']=assignment
  DATA[n]=r;g=GROUPS[u['group']-1];g['topics'].append({k:r[k] for k in ['number','title','kind','hours','outcome']});g['hours']+=2
  note=new_doc();make_note(note,r,u,practical);note.save(DIST/r['download'])

PLAN=dict(name='Server dasturlash (Backend)',hours=180,paras=90,theoryHours=90,practiceHours=90,groups=GROUPS,
 status='Tavsiya etiladigan mualliflik ish rejasi',stack='Node.js 24 LTS · Express 5 · SQLite',author='Ubaydullayev Abdukarim',year='2026/2027')
(ROOT/'authoring/backend-plan.json').write_text(json.dumps(PLAN,ensure_ascii=False,indent=2)+'\n')
(DIST/'backend-data.js').write_text('const backendPlan='+json.dumps(PLAN,ensure_ascii=False)+';\nconst backendGroups=backendPlan.groups;\nconst backendData='+json.dumps(DATA,ensure_ascii=False)+';\n')

work=new_doc();work.add_paragraph('Server dasturlash Backend ish rejasi',style='Title')
work.add_paragraph('Namangan shahar 3 sonli texnikum');work.add_paragraph('Web dasturlash  |  6 modul  |  3 IT  |  2026/2027')
work.add_paragraph('Tayyorlovchi Ubaydullayev Abdukarim')
work.add_paragraph('180 soat  |  90 para  |  90 soat nazariy  |  90 soat amaliy')
work.add_paragraph('Ushbu tavsiya etiladigan mualliflik ish rejasi Frontend modulida o‘tilgan JavaScript bilimlari asosida tuzildi. Rasmiy tasdiqlangan o‘quv dasturi o‘rnida ko‘rsatilmaydi. Har bir tushuntirish darsidan keyin shu bilimga tayangan amaliy mashg‘ulot beriladi.')
work.add_paragraph('Kutiladigan yakuniy natija',style='Heading 1');work.add_paragraph('O‘quvchi HTTP API, SQL baza, parol xeshi, sessiya va resurs egasi tekshiruvini birlashtirgan kichik Backend ilova tuzadi. Uni frontend bilan bog‘laydi, funksional va ruxsat testlarini bajaradi, joylashtirish shartlari va cheklovlarini tushuntiradi.')
work.add_paragraph('Texnologiyalar va tayanch bilimlar',style='Heading 1');work.add_paragraph('Node.js 24 LTS, Express 5, SQLite va node:test. Tayanch: JavaScript funksiya va obyektlari, JSON, async/await, fetch, HTML, CSS hamda Git. Yangi server tushunchalari shu modulda avval izohlanadi.')
work.add_paragraph('Bo‘limlar bo‘yicha taqsimot',style='Heading 1')
tb=work.add_table(rows=1,cols=4)
for c,s in zip(tb.rows[0].cells,['№','Bo‘lim','Para','Soat']):c.text=s
for g in GROUPS:
 for c,s in zip(tb.add_row().cells,[str(g['number']),g['name'],str(len(g['topics'])),str(g['hours'])]):c.text=s
table_style(tb,[.4,5.2,.7,.7])
work.add_paragraph('Baholash va dars tartibi',style='Heading 1');work.add_paragraph('Har mavzuda konspekt, shu konspektga asoslangan slaydlar, 10 ta javobli nazariy savol va 10 ta test mavjud. Har to‘g‘ri test javobi 10 foiz; 60 foizdan kam natija mavzuni qayta o‘qishga yo‘naltiradi. Amaliy ish mezon va hisobot bilan o‘qituvchi tomonidan tekshiriladi; nazariy test bahosi bilan aralashtirilmaydi.')
work.add_paragraph('Mavzular va o‘quv natijalari',style='Heading 1')
for g in GROUPS:
 work.add_paragraph(str(g['number'])+' bo‘lim '+g['name'],style='Heading 2')
 tb=work.add_table(rows=1,cols=5)
 for c,s in zip(tb.rows[0].cells,['№','Mavzu','Shakl','O‘quv natijasi','Soat']):c.text=s
 for t in g['topics']:
  for c,s in zip(tb.add_row().cells,[str(t['number']),t['title'].replace('Amaliy mashg‘ulot — ',''),t['kind'],t['outcome'],'2']):c.text=s
 table_style(tb,[.4,2.35,.9,2.75,.6])
 work.add_paragraph(f'Bo‘lim jami {g["hours"]} soat  |  {len(g["topics"])} para')
work.add_paragraph('Yakuniy loyiha amaliy baholash mezonlari',style='Heading 1')
for x in ['API va CRUD ishlashi 25 ball','Ma’lumot modeli va doimiy saqlash 15 ball','Kirish CSRF va egalik nazorati 25 ball','Testlar va xato holatlari 15 ball','Frontend va foydalanuvchi holatlari 10 ball','README va mustaqil tushuntirish 10 ball']:work.add_paragraph(x,style='List Bullet')
work.add_paragraph('Asosiy manbalar',style='Heading 1')
seen=set()
for values in SOURCES.values():
 for t,url in values:
  if url not in seen:work.add_paragraph(t+' — '+url);seen.add(url)
work.save(MAT/'m6-backend-ish-reja.docx')
with zipfile.ZipFile(MAT/'m6-backend-konspektlar.zip','w',zipfile.ZIP_DEFLATED) as z:
 for n in DATA:z.write(DIST/DATA[n]['download'],Path(DATA[n]['download']).name)

# ZIP contains only reusable source and lockfile, never installed packages or secrets.
project=ROOT/'authoring/backend-project'
with zipfile.ZipFile(MAT/'m6-backend-loyiha.zip','w',zipfile.ZIP_DEFLATED) as z:
 for p in project.rglob('*'):
  if p.is_file() and not any(x in p.relative_to(project).parts for x in ['node_modules','data','.git']) and p.name!='.env':z.write(p,'backend-project/'+str(p.relative_to(project)))
print('Backend built: 90 lessons, 1350 slides, 900 QAs, 900 MCQs, 45 practical assignments, 91 Word files.')
