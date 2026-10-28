import sys,json,html,random,re
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parent))
from frontend_content import L
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
root=Path(__file__).resolve().parents[1];dist=root/'dist';plan=json.loads((root/'authoring/frontend-plan.json').read_text());topics=[t for g in plan['groups'] for t in g['topics']]
esc=lambda v:html.escape(str(v),quote=True)
refs={1:('MDN HTML qo‘llanmasi','https://developer.mozilla.org/en-US/docs/Learn_web_development/Core/Structuring_content'),2:('MDN CSS qo‘llanmasi','https://developer.mozilla.org/en-US/docs/Web/CSS'),3:('Bootstrap 5.3 hujjatlari','https://getbootstrap.com/docs/5.3/getting-started/introduction/'),4:('MDN JavaScript qo‘llanmasi','https://developer.mozilla.org/en-US/docs/Web/JavaScript/Guide'),5:('MDN DOM qo‘llanmasi','https://developer.mozilla.org/en-US/docs/Web/API/Document_Object_Model'),6:('Git rasmiy qo‘llanmasi','https://git-scm.com/docs/gittutorial'),7:('MDN Fetch qo‘llanmasi','https://developer.mozilla.org/en-US/docs/Web/API/Fetch_API/Using_Fetch')}
existing=json.loads((dist/'frontend-resources.js').read_text().split('const frontendResources=',1)[1].split(';\nObject.assign',1)[0]) if (dist/'frontend-resources.js').exists() else {}
resources={int(k):v for k,v in existing.items()}
from frontend_sequence import CHANGED
selected=set(int(x) for x in sys.argv[1].split(',')) if len(sys.argv)>1 else set(range(1,111))
for t in topics:
 if t['number'] not in selected:continue
 n=t['number'];a=L[n];terms=a['terms'];mode=a['mode'];g=next(g for g in plan['groups'] if any(x['number']==n for x in g['topics']))
 objective=f'«{t["title"]}» mavzusida o‘quvchi asosiy tushunchalarni izohlaydi, misolning ishlashini tahlil qiladi va mustaqil vazifani tekshirish mezonlari bilan bajaradi.'
 if mode=='study':objective=f'«{t["title"]}» mavzusining asosiy tushunchalarini o‘z so‘zlari bilan izohlaydi va hayotiy misol orqali farqlaydi.'
 intro='Darsni uch bosqichda o‘rganing: tushunchalarni anglash, namunani bajarib natijani kuzatish, keyin mustaqil topshiriqni o‘zgartirilgan ma’lumot bilan sinash. Tayyor misolni ko‘chirish bilan cheklanmay, har bir muhim yozuv nima vazifa bajarishini tushuntiring.'
 if mode in ['study','files']:intro='Avval tushunchalarni o‘qing, hayotiy misolni tahlil qiling va o‘z so‘zlaringiz bilan izohlang. Kod yozish keyingi bo‘limlarda bosqichma-bosqich boshlanadi.'
 expected=a['expected'] or ('Buyruqlar natijasi haqiqiy loyiha papkasining holatiga bog‘liq. git status yoki tegishli vosita natijasini yozib, kutilgan o‘zgarish bilan solishtiring.' if mode=='git' else 'HTMLni ish maydonida oching. Sahifadagi tuzilma, uslub va hodisalarni kuzating; topshiriqdagi talablar bilan solishtiring. Tashqi media yoki kutubxona talab qilingan bo‘lsa, ularning mavjudligi ham tekshiriladi.')
 solution=('Avval kodni o‘zgartirmasdan ishga tushiring va kutilgan konsol natijasi bilan solishtiring. So‘ng kamida uchta boshqa kirishni tanlang: oddiy, chegara va noto‘g‘ri qiymat. Xato bo‘lsa qiymat turi, amallar tartibi va qaytariladigan natijani navbat bilan tekshiring.' if mode=='js' else 'Avval o‘zingizga tegishli alohida o‘quv papkasida ishlang. Har buyruqdan oldin maqsad va kutilgan holatni yozing; keyin status yoki tarix bilan tasdiqlang. Remote manzil va branchni tekshirmasdan nashr qilmang.' if mode=='git' else 'Avval berilgan misolni ko‘ring, so‘ng vazifadagi o‘zgarishlarni alohida kiriting. Element tuzilishi, CSS qoidasi yoki hodisa qaysi natijani o‘zgartirganini izohlang. Tor va keng ko‘rinishda, klaviatura orqali hamda bo‘sh yoki uzun kontent bilan sinab, kuzatilgan farqlarni qayd eting.')
 solution=a.get('solution',solution)
 if mode in ['study','files']:solution=('Avval tushuntirishni o‘qing, keyin hayotiy misolni o‘z so‘zlaringiz bilan izohlang. Kod yozish talab qilinmaydi.' if mode=='study' else a['teaching'])
 elif mode=='html':solution='Avval konspektdagi teg va atributlarning vazifasini o‘qing. Namunani HTML oynasida ko‘ring, keyin faqat shu va oldingi HTML darslarida o‘tilgan yozuvlarni o‘zgartiring. Sahifa mazmuni va teglar tuzilishini tekshiring. CSS yoki JavaScript kodi yozilmaydi.'
 criteria=['Mavzudagi asosiy tushunchalar to‘g‘ri qo‘llangan.','Topshiriqda sanalgan kontent va amallar bajarilgan.','Oddiy va chegara holatlar uchun kutilgan hamda haqiqiy natija yozilgan.','Xato yoki cheklov aniqlansa sababi va tuzatish yo‘li izohlangan.','Kod yoki hujjat tartibli, nomlar tushunarli va natijani qayta tekshirish mumkin.']
 blocks=[('Darsning maqsadi',objective),('O‘rganish tartibi',intro),('Dars rejasi','\n'.join(f'{i+1}. {k}' for i,(k,v) in enumerate(terms)))]+[(k,v) for k,v in terms]+[('Kod va amaliy namuna',a['code']),('Kutiladigan natija',expected),('Mustaqil topshiriq',a['task']),('Bajarish va tekshirish yo‘li',solution),('Muhim farq va keng tarqalgan xato',a['pitfall']),('Topshirish va baholash','\n'.join(f'{i+1}. {x}' for i,x in enumerate(criteria))+'\nAmaliy ishni baholash: tushuncha va reja 20 ball, bajarilgan yechim 35 ball, natijani sinash 30 ball, izoh va tartib 15 ball. Jami 100 ball. 110-mavzuda loyiha himoyasining vazifada berilgan maxsus mezoni qo‘llanadi. Nazariy test alohida hisoblanadi.')]
 if a.get('teaching'):blocks.insert(3,('Vazifadan oldingi tushuntirish',a['teaching']))
 if a.get('prerequisite'):blocks.insert(2,('Tayanch bilimlar',a['prerequisite']))
 if mode in ['study','files']:
  blocks=[('Hayotiy misol' if h=='Kod va amaliy namuna' else 'Muhokama' if h=='Mustaqil topshiriq' and mode=='study' else h,txt) for h,txt in blocks if h!='Topshirish va baholash']
  criteria=['Asosiy tushuncha o‘z so‘zlari bilan izohlangan.','Misol dars mazmuniga mos.','O‘xshash tushunchalarning farqi aytilgan.','Savolga dalil bilan javob berilgan.','Izoh tushunarli va tartibli.']
 ref=refs[6 if mode=='git' or 79<=n<=82 else 7 if 83<=n<=90 or n==104 else 5 if 69<=n<=78 or 100<=n<=103 else 4 if 47<=n<=68 else 3 if 43<=n<=46 or n==99 else 2 if 19<=n<=42 or 95<=n<=98 else 1]
 web=''
 for h,txt in blocks:
  web+='<h3>'+esc(h)+'</h3>'
  web+=('<pre class="lesson-code"><code>'+esc(txt)+'</code></pre>') if h=='Kod va amaliy namuna' or (h=='Kutiladigan natija' and a['expected']) else ''.join('<p>'+esc(line)+'</p>' for line in txt.split('\n'))
 web+=f'<h3>Qo‘shimcha o‘qish</h3><p><a href="{ref[1]}" target="_blank" rel="noopener">{esc(ref[0])}</a></p><h3>Atamalar</h3><div class="table-scroll"><table><thead><tr><th>Atama</th><th>Izoh</th></tr></thead><tbody>'+''.join('<tr><td>'+esc(k)+'</td><td>'+esc(v)+'</td></tr>' for k,v in terms)+'</tbody></table></div>'
 slides=[{'title':esc(t['title']),'body':[f'5-modul · {g["number"]}-bo‘lim · {n}-mavzu',esc(t['kind'])+' · 2 soat'],'kind':'cover'},{'title':'Dars maqsadi','body':[esc(objective)]},{'title':'Dars yo‘nalishlari','body':[esc(k) for k,v in terms],'kind':'agenda'}]
 slides += [{'title':esc(k),'body':[esc(v)]} for k,v in terms]
 lines=a['code'].split('\n')
 if a.get('teaching'):slides.insert(3,{'title':'Vazifadan oldin','body':[esc(a['teaching'])]})
 # Long HTML lines wrap on the website; split large code blocks into several slides.
 for start in range(0,len(lines),7):slides.append({'title':'Amaliy namuna'+(f' · {start//7+1}' if len(lines)>7 else ''),'body':[],'code':esc('\n'.join(lines[start:start+7]))})
 if mode in ['study','files']:
  for slide in slides:
   if slide.get('code'):slide['title']='Hayotiy misol';slide['body']=[slide.pop('code')]
 slides += [{'title':'Kutiladigan natija','body':[esc(expected)],'kind':'example'},{'title':'Mustaqil vazifa','body':[esc(a['task'])],'kind':'task'},{'title':'Xatoni oldindan ko‘ring','body':[esc(a['pitfall'])],'kind':'review'},{'title':'Natijani tekshirish','body':[esc(x) for x in criteria],'kind':'agenda'},{'title':'Manba va davomiy o‘qish','body':[],'links':[{'title':esc(ref[0]),'url':ref[1]}]}]
 qa=[[esc(k)+' tushunchasi nimani anglatadi?',esc(v)] for k,v in terms]
 qa += [['Namuna qanday natija berishi kutiladi?',esc(expected)],['Ushbu mavzuda qanday xatoga e’tibor berish kerak?',esc(a['pitfall'])],['Mustaqil vazifani qanday bajarish va tekshirish kerak?',esc(a['task']+' '+solution)],['Ish muvaffaqiyatli bajarilganini qanday ko‘rsatasiz?',esc(' '.join(criteria))],[esc(terms[0][0]+' va '+terms[1][0])+' qanday farqlanadi?',esc(terms[0][0]+': '+terms[0][1]+' '+terms[1][0]+': '+terms[1][1])]]
 rng=random.Random(500+n);quiz=[]
 for k,v in terms:
  options=[v]+rng.sample([vv for kk,vv in terms if kk!=k],3);rng.shuffle(options);quiz.append({'q':esc('«'+k+'» uchun to‘g‘ri izohni tanlang.'),'options':[esc(x) for x in options],'correct':options.index(v),'explanation':esc(k+': '+v)})
 for idx,(k,v) in enumerate(terms[:3]):
  options=[kk for kk,vv in terms];options.remove(k);options=rng.sample(options,3)+[k];rng.shuffle(options);quiz.append({'q':esc('Vaziyat: '+v+' Bu qaysi tushunchaga mos?'),'options':[esc(x) for x in options],'correct':options.index(k),'explanation':esc(k+': '+v)})
 for q,right,wrong in [('Misol va topshiriqni tekshirishda qaysi qoida to‘g‘ri?',a['pitfall'],['Faqat kod uzunligi yetarli mezondir.','Natijani ko‘rmasdan barcha test o‘tdi deb yoziladi.','Xato bo‘lsa uning sababini tekshirish shart emas.']),('Dars topshirig‘ining natijasi qanday tasdiqlanadi?',solution,['Faqat mavzu nomini ko‘chirish bilan.','Bitta skrinshot barcha shartni avtomatik isbotlaydi.','Kutilgan natija yozilmasdan barcha holat tasdiqlanadi.'])]:
  opts=[right]+wrong;rng.shuffle(opts);quiz.append({'q':esc(q),'options':[esc(x) for x in opts],'correct':opts.index(right),'explanation':esc(right)})
 resources[n]={**t,'type':'amaliy' if 'amaliy' in t['kind'].lower() else 'nazariy','html':web,'download':f'materials/m5-{n:03d}-konspekt.docx','slides':slides,'qa':qa,'quiz':quiz,'videos':existing.get(str(n),{}).get('videos',[]),'assignment':{'task':a['task'],'code':a['code'],'mode':mode,'expected':expected,'criteria':criteria,'runner':{'functionName':'__frontendExample','tests':[]}}}
 if n<=4:resources[n]['videos']=[]
 if mode=='study':resources[n].pop('assignment',None)
 else:resources[n]['assignment']['teaching']=a.get('teaching','');resources[n]['assignment']['prerequisite']=a.get('prerequisite','')
 d=Document();sec=d.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11);sec.top_margin=sec.bottom_margin=Inches(.7);sec.left_margin=sec.right_margin=Inches(.75)
 for sty in d.styles:
  if sty.type in [1,2]:sty.font.name='Calibri';sty.font.size=Pt(14);sty.font.color.rgb=RGBColor(0,0,0)
  if sty.type==1:
   sty.paragraph_format.line_spacing=1;sty.paragraph_format.space_after=Pt(5)
   if sty.element.pPr is not None:
    for b in list(sty.element.pPr.findall(qn('w:pBdr'))):sty.element.pPr.remove(b)
 for nm in ['Title','Heading 1']:
  st=d.styles[nm];st.font.bold=True;st.paragraph_format.keep_with_next=True;st.paragraph_format.space_before=Pt(8)
 d.add_paragraph('Web interfeysini ishlab chiqish Frontend',style='Title');d.add_paragraph(f'5 modul  {g["number"]} bo‘lim  {n} mavzu  {t["kind"]}  2 soat');d.add_paragraph(re.sub(r'[:;,/()—–]',' ',t['title']),style='Title');d.add_paragraph('Namangan shahar 3-sonli texnikum\nTayyorlovchi Ubaydullayev Abdukarim\n2026–2027 o‘quv yili')
 for h,txt in blocks:
  d.add_heading(re.sub(r'[:;,/()—–]',' ',h),1)
  if h=='Kod va amaliy namuna':
   par=d.add_paragraph(txt);par.paragraph_format.keep_together=True
  else:
   for line in txt.split('\n'):d.add_paragraph(line)
 d.add_heading('Qo‘shimcha o‘qish manbasi',1);d.add_paragraph(ref[0]+'\n'+ref[1]);d.add_heading('Atamalar',1)
 tb=d.add_table(rows=1,cols=2);tb.autofit=False;tb.columns[0].width=Inches(1.7);tb.columns[1].width=Inches(5.3);tb.rows[0].cells[0].text='Atama';tb.rows[0].cells[1].text='Izoh'
 for k,v in terms:cells=tb.add_row().cells;cells[0].text=k;cells[1].text=v
 for i,row in enumerate(tb.rows):
  pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
  if i==0:pr.append(OxmlElement('w:tblHeader'))
  for cell in row.cells:
   cp=cell._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
   for side in ['top','left','bottom','right']:
    z=OxmlElement('w:'+side);z.set(qn('w:val'),'single');z.set(qn('w:sz'),'4');z.set(qn('w:color'),'D9D9D9');b.append(z)
   cp.append(b)
   if i==0:sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'DCE6F1');cp.append(sh)
   for par in cell.paragraphs:
    par.paragraph_format.line_spacing=1;par.paragraph_format.space_before=Pt(4);par.paragraph_format.space_after=Pt(4)
    for run in par.runs:run.font.name='Calibri';run.font.size=Pt(14);run.bold=i==0
 for par in d.paragraphs:
  par.paragraph_format.line_spacing=1
  for run in par.runs:run.font.name='Calibri';run.font.size=Pt(14)
 d.core_properties.author='Ubaydullayev Abdukarim';d.core_properties.title=t['title'];d.save(dist/resources[n]['download'])
for n,data in resources.items():
 if data.get('assignment'):
  data['assignment']['mode']=L[n]['mode']
  if n<=18:data['assignment']['stage']='HTML asoslari' if n>=5 else 'Muharrir va fayllar'
  elif n<=46:data['assignment']['stage']='HTML va CSS; mustaqil JavaScript talab qilinmaydi'
(dist/'frontend-resources.js').write_text('const frontendResources='+json.dumps(resources,ensure_ascii=False)+';\nObject.assign(frontendData,frontendResources);\n')
(root/'authoring/frontend-content.json').write_text(json.dumps(L,ensure_ascii=False,indent=2))
print('Generated',len(resources),'web lessons, DOCX, assignments;',sum(len(v['slides']) for v in resources.values()),'slides;',sum(len(v['quiz']) for v in resources.values()),'MCQs')
