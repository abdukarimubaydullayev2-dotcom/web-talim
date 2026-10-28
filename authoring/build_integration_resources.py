"""Build module 8 without changing prior module content. 45 two-hour lessons."""
import copy,html,json,zipfile
from pathlib import Path
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from integration_units import UNITS
from course_documents import new_doc,table_style,plain_heading
ROOT=Path(__file__).resolve().parents[1];DIST=ROOT/'course';MAT=DIST/'materials';E=lambda x:html.escape(str(x),quote=True)
NAME='Web ilovani integratsiyalash, testlash va joriy etish'
NAMES=['Frontend API va bazani integratsiyalash','Avtomatik unit va integratsion sinovlar','Brauzer sifati qulaylik va xavfsizlik','Reliz CI va joriy etish muhiti','Kuzatish tiklash va yakuniy himoya']
INTRO=[
'Oldingi uch modulda o‘tilgan interfeys, server va ma’lumotlar bazasi bitta oqimga ulanadi. Ushbu moduldagi Vazifalar ilovasi oddiy va tekshirishga qulay misoldir. Har qadamda foydalanuvchi talabi bilan serverdagi natija solishtiriladi.',
'Qo‘lda tekshiruv takrorlanuvchi avtomatik sinovlar bilan to‘ldiriladi. Testlar doimiy o‘quv bazasidan ajratiladi va kutilgan natija aniq yoziladi. Muvaffaqiyatli holat bilan birga noto‘g‘ri kirish ham tekshiriladi.',
'API sinovi brauzerdagi hamma xatoni topmaydi. Shu sababli foydalanuvchi oqimi, klaviatura, ekran o‘lchami va ishonchsiz matnni chiqarish alohida tekshiriladi. O‘lchov yoki nuqson hisobotida muhit va takrorlash qadami ko‘rsatiladi.',
'Sinovdan o‘tgan kodni boshqa muhitda takror ishga tushirish uchun manba holati, sozlama va saqlash talablari qayd etiladi. Reliz jarayoni test bilan bog‘lanadi. Lokal demo tayyorligi bilan ommaviy hostingga chiqarilganlik alohida holatlardir.',
'Ilova ishga tushganidan keyin ham xizmat holati va ma’lumotni saqlash nazorat qilinadi. Log, sog‘liq tekshiruvi, offline backup va tiklash bir-birini to‘ldiradi. Yakuniy himoya bajarilgan ish dalili va hal qilinmagan cheklovlarni o‘z ichiga oladi.'
]
HELP={'js':'Platformadagi Run sof JavaScriptni alohida Workerda bajaradi. console.log va assert(condition, message) mavjud; DOM, node:test, node:sqlite yoki HTTP server bu muhitta ishlamaydi. Namuna tashqi tarmoqqa murojaat qilmaydi.',
'local':'Yuklab olinadigan integration-project katalogida, lokal Node.js 24, terminal va brauzerda bajaring. Playwright va Docker faqat tegishli darsda ixtiyoriy o‘rnatiladi. Platformaning JavaScript oynasi lokal server yoki hostingni yaratmaydi.',
'report':'Bu tahlil va dalil topshirig‘i. Sxema, tekshiruv jadvali va xulosani yozing; matn dastur sifatida bajarilmaydi.'}
SOURCES={
1:[('MDN Fetch','https://developer.mozilla.org/en-US/docs/Web/API/Fetch_API/Using_Fetch'),('MDN CORS','https://developer.mozilla.org/en-US/docs/Web/HTTP/Guides/CORS'),('Node HTTP','https://nodejs.org/docs/latest-v24.x/api/http.html')],
2:[('Node test runner','https://nodejs.org/docs/latest-v24.x/api/test.html'),('Node assert','https://nodejs.org/docs/latest-v24.x/api/assert.html'),('Node SQLite','https://nodejs.org/docs/latest-v24.x/api/sqlite.html')],
3:[('Playwright kirish','https://playwright.dev/docs/intro'),('Playwright lokatorlar','https://playwright.dev/docs/locators'),('MDN accessibility','https://developer.mozilla.org/en-US/docs/Learn_web_development/Core/Accessibility'),('MDN CSP','https://developer.mozilla.org/en-US/docs/Web/HTTP/Guides/CSP')],
4:[('Git reliz taglari','https://git-scm.com/book/en/v2/Git-Basics-Tagging'),('GitHub Node CI','https://docs.github.com/en/actions/tutorials/build-and-test-code/nodejs'),('Docker doimiy saqlash','https://docs.docker.com/get-started/docker-concepts/running-containers/persisting-container-data/')],
5:[('Node HTTP','https://nodejs.org/docs/latest-v24.x/api/http.html'),('SQLite backup','https://sqlite.org/backup.html'),('SQLite tranzaksiyalar','https://sqlite.org/lang_transaction.html')]
}
VIDEOS={13:[dict(id='GYmLXBlTqfE',title='Git va GitHub bilan ishlash — to‘liq kurs',channel='Azizbek Khabibullaev',language='O‘zbekcha',coverage='Git commit va versiya boshqaruvi qismlariga mos. Reliz artefakti va CI alohida konspektda tushuntiriladi; butun kursni qayta bajarish talab qilinmaydi.',task='Videodan commit holatini aniqlash qismini ko‘ring va reliz SHA qiymatini toping.',source='https://www.youtube.com/watch?v=GYmLXBlTqfE',verifiedFrom='https://www.youtube.com/watch?v=GYmLXBlTqfE',checkedOn='2026-10-28')],15:[dict(id='MMP8wCj-5Gg',title='Docker image’lar bilan ishlaymiz — 2-dars',channel='Farkhod Dadajanov',language='O‘zbekcha',coverage='Image tushunchasiga mos eski dars. Joriy Docker o‘rnatish, volume va Node 24 buyruqlari uchun konspektni asos qiling.',task='Image va ishlayotgan container farqini yozing; data katalogi nega volumeda bo‘lishini izohlang.',source='https://www.youtube.com/watch?v=MMP8wCj-5Gg',verifiedFrom='https://www.youtube.com/watch?v=MMP8wCj-5Gg',checkedOn='2026-10-28')]}
CHECKPOINTS=[
('Frontend API va SQLite bilan to‘liq CRUD oqimi','Lokal loyiha ishga tushsin. Vazifa yarating, tugallang va o‘chiring; har amalda Network oynasidagi metod, status va javobni yozing. Bo‘sh title va to‘xtagan server holatini ham kuzating.','npm start\n# Brauzerda http://localhost:3000\n# Networkda POST, PATCH va DELETE javoblarini kuzating.','Uch qatlamda yaratish, yangilash va o‘chirish ishlaydi; xato holati foydalanuvchiga tushunarli chiqadi.'),
('Mustaqil avtomatik test to‘plamini kengaytirish','Paketdagi testlarni o‘qing. 80 belgili title, 81 belgili title, noto‘g‘ri filter va boshqa testdan mustaqillik uchun qo‘shimcha sinov yozing. Bir assertionni buzib xato topilishini isbotlang va qayta tuzating.','npm test\nnode --test test/api.test.mjs','Qo‘shilgan sinovlar ham o‘tadi, ataylab buzilgan assertion esa xato kodi bilan tugaydi.'),
('Brauzer sifati auditi va nuqsonni tuzatish','CRUD oqimini brauzerda tekshiring. Klaviatura, 320 px, 200% zoom va HTMLga o‘xshash title holatini tekshiring. Kamida bitta muammo yoki barcha tekshiruv dalilini yozing; tuzatish bo‘lsa regressiya bilan qayta sinang.','Muhit | Qadam | Kutilgan | Amaldagi | Holat\nKlaviatura | Tab va Enter | boshqariladi | ... | ...\nXSS | title matni | HTML bajarilmaydi | ... | ...','Audit jadvali, takrorlash qadamlari va qayta tekshiruv dalili tayyor bo‘ladi.'),
('Reliz nomzodi va joriy etish dalillari','Toza lokal manba holati va SHAni yozing. Check va testni bajaring. DATA_DIR, port, doimiy disk, HTTPS va qaytish shartlarini docs/release.mdga to‘ldiring. Hosting bo‘lmasa lokal muhitni staging o‘rnida belgilab, HTTPSni tekshirilmadi deb yozing.','npm run check\nnpm test\nnode smoke.mjs http://127.0.0.1:3000','Manba, tekshiruvlar va muhit dalili bog‘lanadi; hostingga chiqarilmagan holat yashirilmaydi.'),
('Yakuniy loyiha namoyishi tiklash va himoya','Qabul mezonlari jadvalini to‘ldiring. Demo yozuv yaratib, serverni to‘xtating, offline backupni yangi katalogga tiklang. Jarayonni, logdagi requestIdni va rollback qarorini ko‘rsating. README, test natijasi va cheklovlarni topshiring.','npm test\n# Server yopilgach:\nnode backup.mjs ./data ./backups/final.sqlite\nnode restore.mjs ./backups/final.sqlite ./restore-demo','Integratsiya, test, tiklash va hujjatlar bitta loyiha himoyasida dalil bilan ko‘rsatiladi.')]
def source_paragraph(doc,values):
 p=doc.add_paragraph()
 for i,(title,url) in enumerate(values):
  if i:p.add_run(' · ')
  link=OxmlElement('w:hyperlink');link.set(qn('r:id'),doc.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True));run=OxmlElement('w:r');t=OxmlElement('w:t');t.text=title;run.append(t);link.append(run);p._p.append(link)
def quiz(u,idx,n):
 qs=[]
 for i,(term,definition) in enumerate(u['concepts']):qs.append(dict(q=term+' haqida qaysi fikr to‘g‘ri?',options=[definition]+[d for j,(t,d) in enumerate(u['concepts']) if i!=j][:3],correct=0,explanation=definition))
 raw=u['case'].split('~');qs.append(dict(q=raw[0],options=raw[1:],correct=0,explanation=raw[1]))
 others=[UNITS[(idx+j)%20] for j in [2,7,11]]
 for question,key in [('Namuna qanday natija berishi kerak?','expected'),('Ushbu mavzuda qaysi xatodan saqlanish kerak?','pitfall'),('O‘tilgan darsdan keyingi amaliy vazifa qaysi?','task')]:qs.append(dict(q=question,options=[u[key]]+[v[key] for v in others],correct=0,explanation=u[key]))
 qs.append(dict(q='Ushbu namuna qaysi muhitda bajariladi?',options=[HELP[u['mode']]]+[v for k,v in HELP.items() if k!=u['mode']]+['Har qanday matn SQL oynasida bajariladi va haqiqiy hostingni avtomatik yaratadi.'],correct=0,explanation=HELP[u['mode']]))
 for i,q in enumerate(qs):
  shift=(i+n)%4;q['options']=q['options'][shift:]+q['options'][:shift];q['correct']=(-shift)%4
  for key in ['q','explanation']:q[key]=E(q[key])
  q['options']=[E(x) for x in q['options']]
 return qs
def make_note(r,u,practical):
 d=new_doc(8);d.styles['Normal'].paragraph_format.space_after=Pt(2)
 d.add_paragraph(plain_heading(r['title']),style='Title');d.add_paragraph(f'8 modul  |  {r["number"]} mavzu  |  2 soat  |  {r["kind"]}')
 d.add_paragraph('Ubaydullayev Abdukarim  |  Namangan shahar 3 sonli texnikum\n3 IT  |  2026/2027')
 for h,s in [('Dars maqsadi',r['outcome']),('Tayanch bilimlar',r['prerequisite']),('Mavzuning mazmuni',INTRO[r['group']-1])]:d.add_paragraph(h,style='Heading 1');d.add_paragraph(s)
 for h,s in u['concepts']:d.add_paragraph(plain_heading(h),style='Heading 2');d.add_paragraph(s)
 d.add_paragraph('Namuna va bajarish muhiti',style='Heading 1');d.add_paragraph(HELP[u['mode']])
 for line in u['code'].splitlines():p=d.add_paragraph(line or ' ');p.paragraph_format.space_after=Pt(0)
 for h,s in [('Kutiladigan natija',u['expected']),('Xatoni tushunish',u['pitfall']),('Amaliy vazifa' if practical else 'Mustahkamlash',u['task'] if practical else 'Namuna oqimini izohlang va kutilgan natijani ayting. Mustaqil ish keyingi amaliy darsda bajariladi.')]:d.add_paragraph(h,style='Heading 1');d.add_paragraph(s)
 if practical:
  for s in r['assignment']['steps']:d.add_paragraph(s,style='List Number')
 d.add_paragraph('Tekshirish mezonlari',style='Heading 2')
 for s in r['criteria']:d.add_paragraph(s,style='List Bullet')
 d.add_paragraph('Atamalar',style='Heading 1')
 for term,definition in u['concepts']:d.add_paragraph(term+' — '+definition.split('. ')[0]+'.')
 d.add_paragraph('Manbalar',style='Heading 1');source_paragraph(d,SOURCES[r['group']]);d.save(DIST/r['download'])
DATA={};GROUPS=[dict(number=i+1,name=x,hours=18,topics=[]) for i,x in enumerate(NAMES)]
def lesson(u,idx,group,practical,checkpoint=False):
 n=len(DATA)+1;title=u['title'] if checkpoint or not practical else 'Amaliy mashg‘ulot — '+u['title'];terms=[t for t,d in u['concepts']]
 outcome=('Topshiriqni bajarib, '+u['expected'][0].lower()+u['expected'][1:] if practical else terms[0]+' va '+terms[1]+' vazifalarini izohlaydi, namuna oqimini va xato holatini tushuntiradi.')
 prereq='5–7-modullardagi HTML, JavaScript, fetch, Node HTTP, SQL va Git asoslari.' if n==1 else f'8-modulning 1–{n-1}-mavzulari. '+('Barcha talab qilingan tushunchalar oldingi nazariy darslarda izohlangan.' if practical else 'Quyidagi tushunchalar shu darsda izohlanadi.')
 criteria=[terms[0]+' misol bilan asoslangan.','Kutilgan va amaldagi natija dalil bilan solishtirilgan.','Tegishli xato yoki cheklov holati qayd etilgan.','Kod yoki hisobotni boshqa o‘quvchi takrorlay oladi.']
 qa=[[t+' nima va qayerda ishlatiladi?',d] for t,d in u['concepts']]+[['Namuna qanday natija beradi?',u['expected']],['Qaysi xatodan saqlanish kerak?',u['pitfall']],['Namuna qaysi muhitda bajariladi?',HELP[u['mode']]],['Amaliy vazifa nimadan iborat?',u['task']],['Natija qanday dalil bilan ko‘rsatiladi?','Test chiqishi yoki kuzatuv jadvalida muhit, qadam, kutilgan va amaldagi natija yoziladi. Tekshirilmagan holat o‘tdi deb belgilanmaydi.']]
 slides=[dict(title=E(title),body=[f'8-modul · {n}-mavzu · 2 soat',E('Amaliy' if practical else 'Nazariy')]),dict(title='Dars maqsadi',body=[E(outcome)]),dict(title='Tayanch va reja',body=[E(prereq),E(' · '.join(terms))])]+[dict(title=E(t),body=[E(s+'.') for s in d.rstrip('.').split('. ')]) for t,d in u['concepts']]+[dict(title='Namuna',body=[],code=E(u['code'])),dict(title='Muhit va natija',body=[E(HELP[u['mode']]),E(u['expected'])]),dict(title='Xatoni tahlil qilish',body=[E(u['pitfall'])]),dict(title='Amaliy vazifa' if practical else 'Mustahkamlash',body=[E(u['task'] if practical else 'Namuna va tushunchalarni izohlang. Mustaqil topshiriq keyingi amaliy darsda beriladi.')]),dict(title='Atamalar',body=[E(t+' — '+d.split('. ')[0]+'.') for t,d in u['concepts']]),dict(title='Tekshirish mezonlari',body=[E(x) for x in criteria]),dict(title='Manbalar',body=[],links=[dict(title=E(t),url=url) for t,url in SOURCES[group]])]
 body=''.join('<h3>'+E(h)+'</h3><p>'+E(s)+'</p>' for h,s in [('Dars maqsadi',outcome),('Tayanch bilimlar',prereq),('Mavzuning mazmuni',INTRO[group-1])]+u['concepts'])
 body+='<h3>Namuna va bajarish muhiti</h3><p>'+E(HELP[u['mode']])+'</p><pre class="lesson-code"><code>'+E(u['code'])+'</code></pre><h3>Kutiladigan natija</h3><p>'+E(u['expected'])+'</p><h3>Xatoni tushunish</h3><p>'+E(u['pitfall'])+'</p><h3>'+('Amaliy vazifa' if practical else 'Mustahkamlash')+'</h3><p>'+E(u['task'] if practical else 'Namuna oqimini izohlang; mustaqil vazifa keyingi amaliy darsda bajariladi.')+'</p><h3>Atamalar</h3>'+''.join('<p><b>'+E(t)+'</b> — '+E(d.split('. ')[0])+'.</p>' for t,d in u['concepts'])+'<h3>Manbalar</h3><ul>'+''.join(f'<li><a href="{url}" target="_blank" rel="noopener">{E(t)}</a></li>' for t,url in SOURCES[group])+'</ul>'
 r=dict(number=n,title=title,group=group,kind='Amaliy' if practical else 'Nazariy',type='amaliy' if practical else 'nazariy',hours=2,unit=idx+1,outcome=outcome,prerequisite=prereq,criteria=criteria,html=body,slides=slides,qa=[[E(q),E(a)] for q,a in qa],quiz=quiz(u,idx,n),download=f'materials/m8-{n:03d}-konspekt.docx',videos=[] if checkpoint else copy.deepcopy(VIDEOS.get(idx+1,[])))
 if practical:r['assignment']=dict(mode=u['mode'],task=u['task'],code=u['code'],expected=u['expected'],draftId=f'm8-r32-{n}',criteria=criteria,steps=['Oldingi nazariy izoh va namuna natijasini o‘qing.','Vazifani belgilangan muhitda bajaring, kutilgan natija bilan solishtiring.','Xato yoki cheklov holatini tekshiring va dalil bilan xulosa yozing.'])
 DATA[n]=r;GROUPS[group-1]['topics'].append({k:r[k] for k in ['number','title','kind','hours','outcome']});make_note(r,u,practical)
for g in range(5):
 for idx in range(g*4,g*4+4):
  for practical in [False,True]:lesson(UNITS[idx],idx,g+1,practical)
 title,task,code,expected=CHECKPOINTS[g];u=copy.deepcopy(UNITS[g*4+3]);u.update(title=title,mode='report' if g in [2,4] else 'local',task=task,code=code,expected=expected)
 # Consolidation uses five concepts already taught in this section; no new prerequisites.
 u['concepts']=[UNITS[g*4+j]['concepts'][0] for j in range(4)]+[UNITS[g*4+3]['concepts'][-1]]
 lesson(u,g*4+3,g+1,True,True)
PLAN=dict(name=NAME,hours=90,paras=45,theoryHours=40,practiceHours=50,groups=GROUPS,author='Ubaydullayev Abdukarim',year='2026/2027',status='Tavsiya etiladigan mualliflik ish rejasi')
(ROOT/'authoring/integration-plan.json').write_text(json.dumps(PLAN,ensure_ascii=False,indent=2)+'\n')
(DIST/'integration-data.js').write_text('const integrationPlan='+json.dumps(PLAN,ensure_ascii=False)+';\nconst integrationGroups=integrationPlan.groups;\nconst integrationData='+json.dumps(DATA,ensure_ascii=False)+';\n')
d=new_doc(8);d.styles['Normal'].paragraph_format.space_after=Pt(2);d.add_paragraph(plain_heading(NAME+' ish rejasi'),style='Title')
for s in ['Namangan shahar 3 sonli texnikum','8 modul  |  3 IT  |  2026/2027','Tayyorlovchi Ubaydullayev Abdukarim','90 soat  |  45 para  |  40 soat nazariy  |  50 soat amaliy']:d.add_paragraph(s)
d.add_paragraph('Tavsiya etiladigan mualliflik reja. 20 ta nazariy va 25 ta amaliy dars. Har nazariy mavzudan keyin tegishli amaliyot, har bo‘lim oxirida umumlashtiruvchi loyiha ishi bor. Rasmiy tasdiqlangan dastur deb belgilanmagan.')
d.add_paragraph('Maqsad va tayanch bilimlar',style='Heading 1');d.add_paragraph('O‘quvchi frontend, API va SQLite bazani birlashtiradi, avtomatik va qo‘lda test qiladi, reliz hamda lokal joriy etish dalilini tayyorlaydi. Tayanch 5–7-modullardagi JavaScript, Node, SQL va Git. Demo Node 24 va brauzerda ishlaydi; Docker va Playwright ixtiyoriy. To‘lov yoki majburiy hosting hisobi talab qilinmaydi.')
d.add_paragraph('Bo‘limlar bo‘yicha soatlar',style='Heading 1');tb=d.add_table(rows=1,cols=4)
for c,s in zip(tb.rows[0].cells,['№','Bo‘lim','Para','Soat']):c.text=s
for g in GROUPS:
 for c,s in zip(tb.add_row().cells,[str(g['number']),g['name'],'9','18']):c.text=s
table_style(tb,[.4,5.1,.75,.75])
d.add_paragraph('Baholash',style='Heading 1');d.add_paragraph('Har mavzuda 10 ta javobli nazariy savol va 10 ta test. Har to‘g‘ri javob 10 foiz; 60 foizdan past natija qayta o‘qishga yo‘naltiradi. Amaliyot o‘qituvchi tomonidan dalil va mezon asosida alohida baholanadi.')
for g in GROUPS:
 d.add_paragraph(str(g['number'])+' bo‘lim '+g['name'],style='Heading 1');tb=d.add_table(rows=1,cols=4)
 for c,s in zip(tb.rows[0].cells,['№','Mavzu','Shakl','Soat']):c.text=s
 for t in g['topics']:
  for c,s in zip(tb.add_row().cells,[str(t['number']),t['title'],t['kind'],'2']):c.text=s
 table_style(tb,[.4,5.2,.85,.55]);d.add_paragraph('Bo‘lim jami 18 soat  |  9 para')
d.add_paragraph('Yakuniy loyiha mezonlari',style='Heading 1')
for s in ['Frontend API va baza integratsiyasi 25 ball','Avtomatik va salbiy testlar 25 ball','Qulaylik xavfsiz matn chiqarish va nuqson dalili 15 ball','Reliz muhit va smoke tekshiruvi 15 ball','Backup tiklash hujjat va himoya 20 ball']:d.add_paragraph(s,style='List Bullet')
d.add_paragraph('Manbalar',style='Heading 1');source_paragraph(d,list(dict.fromkeys(x for values in SOURCES.values() for x in values)));d.save(MAT/'m8-integration-ish-reja.docx')
with zipfile.ZipFile(MAT/'m8-integration-konspektlar.zip','w',zipfile.ZIP_DEFLATED) as z:
 for r in DATA.values():z.write(DIST/r['download'],Path(r['download']).name)
with zipfile.ZipFile(MAT/'m8-integration-loyiha.zip','w',zipfile.ZIP_DEFLATED) as z:
 for p in (ROOT/'authoring/integration-project').rglob('*'):
  if p.is_file() and not any(x in p.parts for x in ['node_modules','data','backups','restore-demo','.git']):z.write(p,'integration-project/'+str(p.relative_to(ROOT/'authoring/integration-project')))
print('45 lessons, 675 slides, 450 QAs and MCQs, 25 assignments, 46 DOCX; theory hours 40, practice hours 50')
