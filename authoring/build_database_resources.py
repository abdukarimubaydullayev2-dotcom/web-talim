"""Build the 140-hour database module without modifying earlier modules."""
import copy,html,json,re,zipfile
from pathlib import Path
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from database_units import UNITS
from database_cases import CASES
from course_documents import new_doc,table_style,plain_heading
ROOT=Path(__file__).resolve().parents[1];DIST=ROOT/'dist';MAT=DIST/'materials'
E=lambda s:html.escape(str(s),quote=True)
NAMES=['Ma’lumotlar modeli va loyihalash','SQLite SQL turlari va cheklovlar','Ma’lumot kiritish o‘zgartirish va qidiruv','JOIN agregatlar va murakkab hisobotlar','Tranzaksiya indeks view va migratsiya','PostgreSQL boshqaruv huquq va backup','Yakuniy kutubxona bazasi loyihasi']
CONTEXT=[
'Bazani qurishdan avval nimani saqlash va qanday qoida bilan bog‘lash aniqlanadi. Ushbu modulda texnikum kutubxonasi izchil misol bo‘ladi. Kitob, a’zo va kitob berish hodisasini ajratish keyingi SQL so‘rovlarini tushunishga tayanch beradi.',
'Model SQL tuzilmasiga aylantiriladi. Tip va cheklov ma’lumotning ma’nosiga mos tanlanadi. Mashqlar SQLite demo muhitida bajariladi; PostgreSQLning qat’iy tiplari va boshqaruv buyruqlari keyingi bo‘limda alohida ko‘riladi.',
'Ma’lumot bilan ishlashda har amalning qamrovi aniq bo‘lishi kerak. Avval kerakli satrlar SELECT bilan kuzatiladi, keyin kiritish yoki o‘zgartirish bajariladi. Natija doimo kutiladigan qiymat bilan solishtiriladi.',
'Hisobot bir jadvaldagi satrlarni chiqarishdan kengroq savolga javob beradi. Bog‘lanish, NULL va guruhlash tanlovi natijadagi satrlar ma’nosini o‘zgartiradi. Har mashqda qaysi obyektlar qolishi kerakligi avval belgilanadi.',
'To‘g‘ri so‘rov bilan birga ma’lumot holatining izchilligi va bajarish xarajati ham muhim. Tranzaksiya o‘zgarish chegarasini, indeks bajarish rejasini, migratsiya esa sxema o‘sishini boshqaradi. Barcha tajribalar demo bazada bajariladi.',
'PostgreSQL alohida server va mijoz orqali ishlaydi. Oldin o‘tilgan SQL bilimlari saqlanadi, biroq schema, rol, izolyatsiya va backup buyruqlari PostgreSQL muhitida tekshiriladi. SQLite oynasi ularni bajarayotgandek ko‘rsatilmaydi.',
'Yakuniy ish model, cheklov, so‘rov, tranzaksiya va tiklash dalillarini birlashtiradi. O‘quvchi faqat skriptni ishga tushirmaydi: har talab qaysi cheklov yoki sinov bilan bajarilishini tushuntiradi. Paketdagi demo ma’lumot haqiqiy shaxsiy ma’lumot emas.'
]
HELP={
'sqlite':'SQLni platformadagi SQLite muharririda bajaring. Har Run yangi demo bazadan boshlanadi; oldingi Run o‘zgarishlari avtomatik qolmaydi. Matn kodi shu brauzerda saqlanadi.',
'report':'Bu model yoki tahlil topshirig‘i. Sxema va yozma hisobot tuzing; matnni SQL sifatida bajarish talab etilmaydi.',
'postgres':'Bu buyruqlar lokal PostgreSQL 18 va psql muhitiga tegishli. Server, demo baza va tegishli huquq kerak; SQLite muharririda bajarilmaydi.',
'node':'Bu JavaScript kodi Node.js 24 va ochiq SQLite ulanishi yoki PostgreSQL drayverini talab qiladi. SQL oynasiga butun Node dasturini joylamang.'}
SOURCES={
1:[('PostgreSQL asoslari','https://www.postgresql.org/docs/18/tutorial.html'),('SQLite tushunchalari','https://sqlite.org/whentouse.html')],
2:[('SQLite SQL','https://sqlite.org/lang.html'),('SQLite tiplari','https://sqlite.org/datatype3.html')],
3:[('SQLite SELECT','https://sqlite.org/lang_select.html'),('SQLite ifodalar','https://sqlite.org/lang_expr.html')],
4:[('SQLite SELECT va JOIN','https://sqlite.org/lang_select.html'),('SQLite oynali funksiyalar','https://sqlite.org/windowfunctions.html')],
5:[('SQLite tranzaksiyalar','https://sqlite.org/lang_transaction.html'),('SQLite so‘rov rejalari','https://sqlite.org/eqp.html')],
6:[('PostgreSQL 18 qo‘llanmasi','https://www.postgresql.org/docs/18/tutorial.html'),('PostgreSQL backup','https://www.postgresql.org/docs/18/backup-dump.html')],
7:[('SQLite foreign key','https://sqlite.org/foreignkeys.html'),('Node SQLite','https://nodejs.org/api/sqlite.html')]
}
VIDEOS={}
def video(units,ident,title,channel,note,evidence=None):
 for u in units:VIDEOS.setdefault(u,[]).append(dict(id=ident,title=title,channel=channel,language='O‘zbekcha',coverage=note,task='Videodagi mavzuga tegishli qismni konspekt bilan solishtiring. Hozirgi darsda o‘tilgan tushunchani izohlang; keyingi mavzu buyruqlari hozircha majburiy emas.',source='https://www.youtube.com/watch?v='+ident,verifiedFrom=evidence or 'https://www.youtube.com/watch?v='+ident,checkedOn='2026-10-28'))
video([1],'ICFe2W5Gyss','SQL — ma’lumotlar ombori haqida tushuncha','Bexruz Xolmuminov','Ma’lumotlar bazasi va SQLga kirish. PostgreSQL kursidan; hozircha umumiy tushunchalar qismi kerak.')
video([6],'2L4dye3x1Tk','SQL — DB Browser for SQLite','Sanjar Suvonov','SQLite vositasi bilan tanishishga mos. Android kursidagi dars; Android ilova tuzish bu mavzuda talab etilmaydi.','https://t.me/s/mobileracademy?after=117')
video([9,17],'VXSzVtV-Lpk','PostgreSQL — FOREIGN KEY, RELATIONSHIP va JOIN','Bexruz Xolmuminov','Bog‘lanish va tashqi kalitlar haqida. PostgreSQL sintaksisi ishlatiladi; SQLite mashq uchun konspektni asos qiling. Video barcha JOIN turlarini to‘liq qamraydi deb hisoblanmaydi.')
video([10],'9HFRpkPh-gA','SQL — ALTER TABLE va ustun qo‘shish','ITOne Uz','Ustun qo‘shishga aloqador dars. Muallif MySQL muhitida ishlaydi; SQLite sintaksisi konspektda alohida berilgan.','https://t.me/s/itone_uz?before=331')
video([11],'oL7nFSLEYtg','PostgreSQL — INSERT INTO','Bexruz Xolmuminov','INSERT orqali yozuv kiritish. PostgreSQL va SQLite farqlari bo‘lsa konspektdagi muhitga amal qiling.')
video([12],'oekR8yp4CtE','SQL — SELECT, 1-qism','ITOne Uz','SELECT tushunchasi uchun. Videoda MySQL; mashqda esa SQLite ishlatiladi.','https://t.me/s/itone_uz?before=316')
video([12],'_KJzMkffd_k','SQL — SELECT va WHERE, 2-qism','ITOne Uz','WHERE bilan filtrlashga aloqador dars. MySQL va SQLite farqlarida konspektni asos qiling.','https://t.me/s/itone_uz?before=316')
video([17],'FX4AFqP9Jow','SQL — JOIN va jadvallarni bog‘lash','ITOne Uz','JOIN tushunchasi uchun. Videodagi MySQL muhiti ushbu SQLite mashq muhitidan farq qiladi.','https://t.me/s/itone_uz?before=331')
video([18],'DetlSJ0X5hQ','PostgreSQL — LEFT JOIN','Bexruz Xolmuminov','LEFT JOINga oid dars. Nol berishlarni sanash va COUNT farqi konspektda qo‘shimcha tushuntirilgan.')
video([26],'OMVclONSfyA','PostgreSQLni kompyuterga o‘rnatish','Bexruz Xolmuminov','O‘rnatish jarayoniga aloqador dars. Video eski relizga tegishli bo‘lishi mumkin; kurs PostgreSQL 18 hujjatlariga tayanadi, yangi o‘rnatgich ko‘rinishi farqlanishi mumkin.')

def quiz_for(u,index,n):
 qs=[]
 for i,(term,defn) in enumerate(u['concepts']):qs.append(dict(q=term+' bo‘yicha qaysi izoh to‘g‘ri?',options=[defn]+[x[1] for j,x in enumerate(u['concepts']) if i!=j][:3],correct=0,explanation=defn))
 for raw in CASES[index]:
  q,*options=raw.split('~');qs.append(dict(q=q,options=options,correct=0,explanation=options[0]+'.'))
 others=[UNITS[(index+j)%35] for j in [1,8,14]]
 qs += [dict(q='Konspektdagi namuna uchun qaysi natija kutiladi?',options=[u['expected']]+[v['expected'] for v in others],correct=0,explanation=u['expected']),dict(q='Ushbu darsdagi mashqda qaysi xatoga e’tibor beriladi?',options=[u['pitfall']]+[v['pitfall'] for v in others],correct=0,explanation=u['pitfall']),dict(q='Ushbu namuna qaysi muhit va tartibda bajariladi?',options=[HELP[u['mode']]]+[s for k,s in HELP.items() if k!=u['mode']],correct=0,explanation=HELP[u['mode']])]
 for i,q in enumerate(qs):
  shift=(n+i)%4;q['options']=q['options'][shift:]+q['options'][:shift];q['correct']=(-shift)%4
  q['q']=E(q['q']);q['options']=[E(x) for x in q['options']];q['explanation']=E(q['explanation'])
 return qs

def source_paragraph(doc,values):
 p=doc.add_paragraph()
 for i,(title,url) in enumerate(values):
  if i:p.add_run(' · ')
  link=OxmlElement('w:hyperlink');link.set(qn('r:id'),doc.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True))
  run=OxmlElement('w:r');text=OxmlElement('w:t');text.text=title;run.append(text);link.append(run);p._p.append(link)

def note(r,u,practical):
 d=new_doc(7);d.styles['Normal'].paragraph_format.space_after=Pt(2)
 d.add_paragraph(plain_heading(r['title']),style='Title');d.add_paragraph(f'7 modul  |  {r["number"]} mavzu  |  2 soat  |  {r["kind"]}')
 d.add_paragraph('Ubaydullayev Abdukarim  |  Namangan shahar 3 sonli texnikum\n3 IT  |  2026/2027')
 d.add_paragraph('Dars maqsadi',style='Heading 1');d.add_paragraph(r['outcome'])
 d.add_paragraph('Tayanch bilimlar',style='Heading 2');d.add_paragraph(r['prerequisite'])
 d.add_paragraph('Mavzuning mazmuni',style='Heading 1');d.add_paragraph('Oldingi nazariy darsdagi tushunchalarni misolda qo‘llang. Natija, cheklov va xato holatini quyidagi tartibda tekshiring.' if practical else CONTEXT[u['group']-1])
 for title,defn in u['concepts']:d.add_paragraph(plain_heading(title),style='Heading 2');d.add_paragraph(defn)
 d.add_paragraph('Namuna va bajarish muhiti',style='Heading 1');d.add_paragraph(HELP[u['mode']])
 for line in u['code'].splitlines():
  p=d.add_paragraph(line or ' ');p.paragraph_format.space_after=Pt(0)
 d.add_paragraph('Kutiladigan natija',style='Heading 2');d.add_paragraph(u['expected'])
 d.add_paragraph('Xatoni tushunish',style='Heading 2');d.add_paragraph(u['pitfall'])
 d.add_paragraph('Amaliy ish' if practical else 'Mustahkamlash',style='Heading 1');d.add_paragraph(u['task'] if practical else 'Asosiy tushunchalarni misol bilan izohlang. Namunaning kirish va natijasini ajrating; mustaqil vazifa keyingi amaliy darsda beriladi.')
 if practical:
  for s in r['assignment']['steps']:d.add_paragraph(s,style='List Number')
 d.add_paragraph('Tekshirish mezonlari',style='Heading 2')
 for s in r['criteria']:d.add_paragraph(s,style='List Bullet')
 d.add_paragraph('Atamalar',style='Heading 1')
 for t,defn in u['concepts']:d.add_paragraph(t+' — '+defn.split('. ')[0]+'.')
 d.add_paragraph('Manbalar',style='Heading 1');source_paragraph(d,SOURCES[u['group']]);d.save(DIST/r['download'])

DATA={};GROUPS=[dict(number=i+1,name=name,hours=20,topics=[]) for i,name in enumerate(NAMES)]
for index,u in enumerate(UNITS):
 for practical in [False,True]:
  n=index*2+1+int(practical);terms=[t for t,d in u['concepts']];title=('Amaliy mashg‘ulot — ' if practical else '')+u['title']
  outcome=(f'{terms[0]} va {terms[1]} vazifalarini izohlaydi; {terms[2]} qo‘llanishini misolda ajratadi.' if not practical else f'{u["title"]} mavzusi bo‘yicha amaliy natija oladi va uning to‘g‘riligini dalil bilan tekshiradi.')
  prereq='6-modulda o‘tilgan SQLite CRUD va JavaScript asoslari.' if n==1 else (f'{n-1}-mavzudagi nazariy tushuntirish va undan oldingi darslar.' if practical else f'7-modulning 1–{n-1}-mavzulari. Yangi tushunchalar ushbu darsda izohlanadi.')
  criteria=[terms[0]+' vazifasi misol bilan tushuntirilgan.','Namuna muhiti va natijasi to‘g‘ri ajratilgan.','Tegishli xato yoki chegaraviy holat tekshirilgan.','Natija va xulosa yozilgan.']
  qa=[[t+' nima va bu mavzuda qanday ishlatiladi?',d] for t,d in u['concepts']]+[['Namuna qanday natija beradi?',u['expected']],['Qaysi xatoga e’tibor beriladi?',u['pitfall']],['Qaysi muhitda bajariladi?',HELP[u['mode']]],['Amaliy ishda nima bajariladi?',u['task']],['Tekshirish uchun qanday dalil kerak?','Kutilgan natija bilan olingan satr va qiymatlar solishtiriladi; tegishli mavzuda noto‘g‘ri kirish ham tekshiriladi. Sxema ishida kalit va bog‘lanishlar asoslanadi.']]
  slides=[dict(title=E(title),body=[f'7-modul · {n}-mavzu · 2 soat',E('Amaliy' if practical else 'Nazariy')]),dict(title='Dars maqsadi',body=[E(outcome)]),dict(title='Tushunchalar',body=[E(' · '.join(terms))],kind='agenda')]
  slides += [dict(title=E(t),body=[E(x+'.') for x in d.rstrip('.').split('. ')],kind='example') for t,d in u['concepts']]
  slides += [dict(title='Namuna',body=[],code=E(u['code']),kind='example'),dict(title='Natija va muhit',body=[E(u['expected']),E(HELP[u['mode']])],kind='solution'),dict(title='Xatoni tahlil qilish',body=[E(u['pitfall'])],kind='review'),dict(title='Amaliy vazifa' if practical else 'Mustahkamlash',body=[E(u['task'] if practical else 'Tushunchalarni misolda izohlang. Mustaqil bajarish keyingi amaliy darsda beriladi.')],kind='task'),dict(title='Atamalar',body=[E(t+' — '+d.split('. ')[0]+'.') for t,d in u['concepts']],kind='terms'),dict(title='Tekshirish mezonlari',body=[E(s) for s in criteria],kind='review'),dict(title='Manbalar',body=[],links=[dict(title=E(t),url=url) for t,url in SOURCES[u['group']]])]
  body='<h3>Dars maqsadi</h3><p>'+E(outcome)+'</p><h3>Tayanch bilimlar</h3><p>'+E(prereq)+'</p><h3>Mavzuning mazmuni</h3><p>'+E(CONTEXT[u['group']-1])+'</p>'
  body+=''.join('<h3>'+E(t)+'</h3><p>'+E(d)+'</p>' for t,d in u['concepts'])
  body+='<h3>Namuna va bajarish muhiti</h3><p>'+E(HELP[u['mode']])+'</p><pre class="lesson-code"><code>'+E(u['code'])+'</code></pre><h3>Kutiladigan natija</h3><p>'+E(u['expected'])+'</p><h3>Xatoni tushunish</h3><p>'+E(u['pitfall'])+'</p>'
  body+='<h3>'+('Amaliy vazifa' if practical else 'Mustahkamlash')+'</h3><p>'+E(u['task'] if practical else 'Tushunchalarni o‘z so‘zingiz bilan izohlang va namunani kuzating. Mustaqil vazifa keyingi amaliy darsda bajariladi.')+'</p><h3>Atamalar</h3>'
  body+=''.join('<p><b>'+E(t)+'</b> — '+E(d.split('. ')[0])+'.</p>' for t,d in u['concepts'])
  body+='<h3>Manbalar</h3><ul>'+''.join(f'<li><a href="{url}" target="_blank" rel="noopener">{E(t)}</a></li>' for t,url in SOURCES[u['group']])+'</ul>'
  r=dict(number=n,title=title,kind='Amaliy' if practical else 'Nazariy',type='amaliy' if practical else 'nazariy',hours=2,unit=index+1,outcome=outcome,prerequisite=prereq,criteria=criteria,html=body,slides=slides,qa=[[E(q),E(a)] for q,a in qa],quiz=quiz_for(u,index,n),download=f'materials/m7-{n:03d}-konspekt.docx',videos=copy.deepcopy(VIDEOS.get(index+1,[])))
  if practical:r['assignment']=dict(mode=u['mode'],task=u['task'],code=u['code'],expected=u['expected'],useFixture=u['fixture'],draftId=f'm7-r31-{n}',unit=index+1,criteria=criteria,steps=['Oldingi nazariy dars va namuna izohini o‘qing.','Topshiriqni berilgan muhitda bajaring va natijani solishtiring.','Tegishli xato holatini tekshirib, xulosani yozing.'])
  DATA[n]=r;GROUPS[u['group']-1]['topics'].append({k:r[k] for k in ['number','title','kind','hours','outcome']});note(r,u,practical)

PLAN=dict(name='Ma’lumotlar bazasi bilan ishlash',hours=140,paras=70,theoryHours=70,practiceHours=70,groups=GROUPS,author='Ubaydullayev Abdukarim',year='2026/2027',status='Tavsiya etiladigan mualliflik ish rejasi')
fixture=(ROOT/'authoring/database_schema.sql').read_text()+'\n'+(ROOT/'authoring/database_seed.sql').read_text()
(ROOT/'authoring/database-plan.json').write_text(json.dumps(PLAN,ensure_ascii=False,indent=2)+'\n')
(DIST/'database-data.js').write_text('const databasePlan='+json.dumps(PLAN,ensure_ascii=False)+';\nconst databaseGroups=databasePlan.groups;\nconst databaseFixture='+json.dumps(fixture,ensure_ascii=False)+';\nconst databaseData='+json.dumps(DATA,ensure_ascii=False)+';\n')

d=new_doc(7);d.add_paragraph('Ma’lumotlar bazasi bilan ishlash ish rejasi',style='Title')
for s in ['Namangan shahar 3 sonli texnikum','7 modul  |  3 IT  |  2026/2027','Tayyorlovchi Ubaydullayev Abdukarim','140 soat  |  70 para  |  70 soat nazariy  |  70 soat amaliy']:d.add_paragraph(s)
d.add_paragraph('Ushbu tavsiya etiladigan mualliflik reja 6-moduldagi SQLite va Backend bilimlarini chuqurlashtiradi. Rasmiy tasdiqlangan dastur sifatida belgilanmagan. Har bir nazariy darsdan keyin shu mavzuga mos amaliy mashg‘ulot keladi.')
d.add_paragraph('Kutiladigan natija',style='Heading 1');d.add_paragraph('O‘quvchi relatsion model va ER diagramma tuzadi, SQL bilan ma’lumotni boshqaradi, JOIN va agregat hisobot oladi. Tranzaksiya, indeks, ruxsat va backup vazifalarini farqlaydi, kutubxona demo bazasini sinab tiklash dalilini ko‘rsatadi.')
d.add_paragraph('Muhit va tayanch bilimlar',style='Heading 1');d.add_paragraph('SQLite brauzer muharriri, DB Browser yoki sqlite3, Node.js 24 uchun test fayli, lokal PostgreSQL 18 va psql. PostgreSQLga xos amallar SQLite muharririda bajarilmaydi. Tayanch: 6-moduldagi CRUD, Node va parametrli SQLite so‘rovlari.')
d.add_paragraph('Bo‘limlar bo‘yicha soatlar',style='Heading 1');tb=d.add_table(rows=1,cols=4)
for c,s in zip(tb.rows[0].cells,['№','Bo‘lim','Para','Soat']):c.text=s
for g in GROUPS:
 for c,s in zip(tb.add_row().cells,[str(g['number']),g['name'],'10','20']):c.text=s
table_style(tb,[.4,5.1,.75,.75])
d.add_paragraph('Baholash',style='Heading 1');d.add_paragraph('Har mavzu 10 ta javobli savol va 10 ta test bilan yakunlanadi. Har to‘g‘ri javob 10 foiz; 60 foizdan past natija qayta o‘qishga yo‘naltiradi. Amaliy ish o‘qituvchi tomonidan mezon va hisobot asosida alohida baholanadi.')
for g in GROUPS:
 d.add_paragraph(str(g['number'])+' bo‘lim '+g['name'],style='Heading 1');tb=d.add_table(rows=1,cols=5)
 for c,s in zip(tb.rows[0].cells,['№','Mavzu','Shakl','O‘quv natijasi','Soat']):c.text=s
 for t in g['topics']:
  for c,s in zip(tb.add_row().cells,[str(t['number']),t['title'].replace('Amaliy mashg‘ulot — ',''),t['kind'],t['outcome'],'2']):c.text=s
 table_style(tb,[.4,2.35,.9,2.75,.6]);d.add_paragraph('Bo‘lim jami 20 soat  |  10 para')
d.add_paragraph('Yakuniy loyiha amaliy bahosi',style='Heading 1')
for s in ['ER model va ustunlar lug‘ati 20 ball','Sxema kalitlar va cheklovlar 25 ball','To‘g‘ri hisobot so‘rovlari 20 ball','Tranzaksiya va salbiy sinovlar 15 ball','Backup va tiklash dalili 10 ball','Hujjat va mustaqil tushuntirish 10 ball']:d.add_paragraph(s,style='List Bullet')
d.add_paragraph('Manbalar',style='Heading 1');seen=set();unique_sources=[]
for values in SOURCES.values():
 for title,url in values:
  if url not in seen:unique_sources.append((title,url));seen.add(url)
source_paragraph(d,unique_sources)
d.save(MAT/'m7-database-ish-reja.docx')
with zipfile.ZipFile(MAT/'m7-database-konspektlar.zip','w',zipfile.ZIP_DEFLATED) as z:
 for n in DATA:z.write(DIST/DATA[n]['download'],Path(DATA[n]['download']).name)
with zipfile.ZipFile(MAT/'m7-database-loyiha.zip','w',zipfile.ZIP_DEFLATED) as z:
 for p in (ROOT/'authoring/database-project').rglob('*'):
  if p.is_file():z.write(p,'database-project/'+str(p.relative_to(ROOT/'authoring/database-project')))
print('70 lessons, 1050 slides, 700 QAs, 700 MCQs, 35 practical assignments, 71 DOCX files; videos:',sum(bool(r['videos']) for r in DATA.values()))
