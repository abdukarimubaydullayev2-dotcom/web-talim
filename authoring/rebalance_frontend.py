"""Build the agreed 16-hour HTML / 56-hour JavaScript curriculum from retained source."""
from pathlib import Path
from copy import deepcopy
import json,re,html,random,sys
from bs4 import BeautifulSoup
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1];DIST=ROOT/'dist';BASE=json.loads((ROOT/'authoring/frontend-before-rebalance.json').read_text());old={int(k):v for k,v in BASE['resources'].items()};content={int(k):v for k,v in BASE['content'].items()}
html_units=[[5,6],[7,8],[9,10],[11,12],[13,14],[15,16],[17],[18]]
html_titles=['HTML hujjati tuzilishi, teglar, elementlar va atributlar','Sarlavhalar, paragraflar, matnni formatlash va maxsus belgilar','Giperhavolalar, rasmlar va fayl yo‘llari','Tartibli va tartibsiz ro‘yxatlar, HTML jadvallari','Semantik HTML va sahifa navigatsiyasi','HTML formalar va foydalanuvchi kiritish elementlari','Audio, video va multimedia elementlari','Head, title, meta va HTML sahifani yakuniy tekshirish']
js_order=[47,48,49,50,51,52,'practice-values',53,54,55,'practice-branches',56,57,58,'practice-loops',59,60,61,'practice-arrays',62,63,64,'practice-functions',65,66,67,68,'practice-project']
units={**{n:[n] for n in range(1,5)},**{5+i:v for i,v in enumerate(html_units)},**{n-6:[n] for n in range(19,47)},**{41+i:[v] for i,v in enumerate(js_order)},**{n:[n] for n in range(69,111)}}
number_map={v:k for k,vs in units.items() for v in vs if isinstance(v,int)}
def renumber_text(s):
 def repl(m):
  a=int(m[1]);b=int(m[2]) if m[2] else None
  return str(number_map.get(a,a))+(('–'+str(number_map.get(b,b))) if b is not None else '')+'-mavzu'
 return re.sub(r'(\d+)(?:[–−-](\d+))?-mavzu',repl,s)
def remap(value):
 if isinstance(value,str):return renumber_text(value)
 if isinstance(value,list):return [remap(v) for v in value]
 if isinstance(value,dict):return {k:remap(v) for k,v in value.items()}
 return value
P={}
def practice(key,title,terms,teaching,code,expected,task,pitfall,checks):
 P[key]={'title':title,'terms':terms,'teaching':teaching,'code':code,'expected':expected,'task':task,'pitfall':pitfall,'checks':checks}
practice('practice-values','Amaliy mustahkamlash: o‘zgaruvchilar, turlar, operatorlar va kiritish-chiqarish',
 [('Bog‘lanish','let qayta qiymat beriladigan, const esa qayta bog‘lanmaydigan nomni ifodalaydi.'),('Tur','Son va satr farqlanadi: arifmetik amal oldidan satrning son ko‘rinishi Number bilan olinadi.'),('O‘zlashtirish','Qiymatni nomga yozishdir; = qiymat beradi, === esa tenglikni tekshiradi.'),('Kuzatish','console.log qiymat yoki hisob natijasini konsolda ko‘rsatadi; u sahifa elementini o‘zgartirmaydi.'),('Kirish','prompt javobi satr yoki null bo‘lishi mumkin. Number.isFinite aylantirilgan qiymatni kuzatishga yordam beradi.')],
 'Avval 43–46-mavzulardagi let, const, ma’lumot turi va operatorlarni eslang. Narx va miqdorni sonlarda saqlang, ko‘paytirish natijasini boshqa nomga yozing. So‘ng xuddi shu qiymatlarni prompt orqali berib Number bilan aylantiring. Hozir shart, sikl, massiv yoki DOM talab qilinmaydi; noto‘g‘ri qiymat kelganda chiqadigan natijani kuzatishning o‘zi kifoya.',
 'const narx = 12000;\nlet miqdor = 3;\nlet jami = narx * miqdor;\nconsole.log(jami);\nmiqdor = 5;\njami = narx * miqdor;\nconsole.log(jami);','36000\n60000',
 'Daftar narxi va soni asosida jami to‘lovni hisoblang. Avval tayyor sonlar, keyin prompt va Number bilan ishlang. Uch xil miqdor uchun kutilgan va haqiqiy natijani yozing. if yoki funksiya yaratish shart emas.',
 '"12000" + "3" satrlarni ulaydi; son hisoblash uchun qiymat turi oldin tekshiriladi.',
 [('12000 × 3','36000'),('12000 × 0','0'),('12000 × 5','60000')])
practice('practice-branches','Amaliy mustahkamlash: if, else va switch bilan qaror chiqarish',
 [('Shart','Boolean natija asosida qaysi blok ishlashini tanlaydi.'),('Chegara','Qaror o‘zgaradigan qiymat; masalan 60 ballning o‘zi ham sinovga kiradi.'),('Yaroqli oraliq','Kirishning qabul qilinadigan qiymat chegaralaridir; 0–100 ball misolida manfiy yoki 101 qiymat yaroqsiz.'),('Ketma-ket tekshirish','if va else if yuqoridan pastga tekshiriladi; birinchi mos tarmoq bajariladi.'),('switch','Aniq variantlar uchun case va default bilan tanlash; break keyingi casega o‘tishni to‘xtatadi.')],
 '48–50-mavzulardagi if/else va switchdan foydalaning. Ballni baholashda avval yaroqsiz oraliqni, keyin o‘tish chegarasini tekshiring. if va switchni bitta usulga majburlamang: ball oraliqlari ifga, amal nomlari switchga qulay. Sikl yoki massiv hali kerak emas.',
 'const ball = 60;\nif (ball < 0 || ball > 100) { console.log("Yaroqsiz ball"); }\nelse if (ball >= 60) { console.log("O‘tdi"); }\nelse { console.log("Qayta o‘qish"); }','O‘tdi',
 'Ball bo‘yicha qaror dasturini yozing. −1, 0, 59, 60, 100 va 101 uchun natijani oldindan yozib, kod bilan solishtiring. Keyin +, −, * va / amallarini switch bilan tanlang; nolga bo‘lishni alohida tekshiring.',
 'Umumiy shartni birinchi qo‘yish maxsus holatni yopib qo‘yishi mumkin; chegaralarni ikki tomondan sinang.',
 [('59','Qayta o‘qish'),('60','O‘tdi'),('101','Yaroqsiz ball')])
practice('practice-loops','Amaliy mustahkamlash: sikllar, hisoblagich va yig‘indi masalalari',
 [('Hisoblagich','Har aylanishda belgilangan qadam bilan o‘zgaradigan qiymat.'),('Akkumulyator','Sikl davomida umumiy natijani yig‘adigan o‘zgaruvchi.'),('Boshlang‘ich qiymat','Yig‘indi uchun 0 tanlanadi; noto‘g‘ri boshlanish butun natijani o‘zgartiradi.'),('Tugash sharti','Siklning qachon to‘xtashini belgilaydi; hisoblagich shartga yaqinlashishi kerak.'),('Solishtirish','Bir masalani for va while bilan yechib, bir xil kirishda natijalarini tekshirish.')],
 '52–54-mavzulardagi for, while, do...while, break va continuega tayaning. Avval nning kichik qiymatini tanlab, har aylanishda i va summa qanday o‘zgarganini qog‘ozda yozing. Keyin kodni bajaring. Massiv, funksiya va sahifa elementlari bilan ishlash bu mashqning talabi emas.',
 'const n = 5;\nlet summa = 0;\nfor (let i = 1; i <= n; i++) { summa += i; }\nconsole.log(summa);\nlet juftlar = 0;\nfor (let i = 1; i <= n; i++) { if (i % 2 === 0) juftlar++; }\nconsole.log(juftlar);','15\n2',
 '1 dan n gacha yig‘indi va juft sonlar miqdorini hisoblang. for yechimini whilega o‘giring. n=1, n=5 va n=10 holatlarini tekshiring; har bir siklning tugash shartini tushuntiring.',
 'i < n va i <= n bir xil emas. Hisoblagich yangilanmasa sikl tugamasligi mumkin.',
 [('n=1: yig‘indi / juftlar','1 / 0'),('n=5','15 / 2'),('n=10','55 / 5')])
practice('practice-arrays','Amaliy mustahkamlash: massivlar va ma’lumotlarni saralash',
 [('Indeks','Massivdagi o‘rin 0 dan boshlanadi; oxirgi odatiy indeks length − 1.'),('Massivni ko‘rish','Sikl bilan har elementga indeks orqali murojaat qilish.'),('Filtrlash','Shartga mos elementlarni ajratish; bu darsda oddiy for va push bilan bajarish mumkin.'),('Yangi massiv','Natijani alohida massivga yig‘ish asl ro‘yxatni saqlab qoladi.'),('Bo‘sh natija','Hech bir element shartga mos kelmagan holat ham tekshiriladi.')],
 '56–58-mavzularda massiv va metodlar ko‘rildi. Avval for bilan elementlarni o‘qing, if bilan shartni tekshiring, push bilan yangi ro‘yxatga qo‘shing. map/filter uchun oldin tushuntirilgan callback qoliplarini kuzatish mumkin; yangi funksiya yozish bu mashqda majburiy emas.',
 'const ballar = [40, 60, 80];\nconst otganlar = [];\nfor (let i = 0; i < ballar.length; i++) {\n if (ballar[i] >= 60) otganlar.push(ballar[i]);\n}\nconsole.log(otganlar);\nconsole.log(otganlar.length);','[60,80]\n2',
 'Ballar ro‘yxatidan 60 va undan yuqori natijalarni alohida massivga yig‘ing. Asl ro‘yxat o‘zgarmasin. Bo‘sh ro‘yxat, hech kim o‘tmagan ro‘yxat va barcha o‘tgan ro‘yxatni sinang.',
 'i <= massiv.length odatda mavjud bo‘lmagan indeksni ham o‘qiydi; i < massiv.length ishlatiladi.',
 [('[40,60,80]','[60,80]'),('[]','[]'),('[10,20]','[]')])
practice('practice-functions','Amaliy mustahkamlash: funksiyalar, parametrlar va qaytariladigan qiymatlar',
 [('Parametr','Funksiya ta’rifidagi kirish nomi; argument chaqiruvdagi haqiqiy qiymat.'),('return','Natijani chaqiruv joyiga qaytaradi va funksiyani tugatadi.'),('Standart qiymat','Argument berilmaganda qo‘llanadigan boshlang‘ich parametr qiymati.'),('Sof hisoblash','Bir xil kirishga bir xil hisob natijasi beruvchi, tashqi holatni o‘zgartirmaydigan funksiya.'),('Qayta foydalanish','Bitta hisoblashni turli argumentlar bilan chaqirish; natijani alohida chiqarish.')],
 '60–62-mavzulardagi e’lon, parametr, return va arrow yozuvini eslang. Hisobni funksiya ichida bajaring va return qiling; console.logni tashqarida ishlating. Avval oddiy function yozuvini sinang, so‘ng xuddi shu hisobni arrow bilan ifodalang. Obyekt va DOM hali vazifa emas.',
 'function tolov(narx, son = 1) { return narx * son; }\nconsole.log(tolov(5000));\nconsole.log(tolov(5000, 3));\nconst kvadrat = x => x * x;\nconsole.log(kvadrat(4));','5000\n15000\n16',
 'Narx va miqdorni qabul qilib jami to‘lovni qaytaradigan funksiya yozing. Miqdor berilmasa 1 bo‘lsin. Uch xil chaqiruv bilan tekshiring. Oddiy functionni arrowga o‘giring va natijalar bir xil ekanini tekshiring.',
 'console.log natijani ko‘rsatadi, ammo funksiyadan qiymat qaytarmaydi. Returnni ko‘rsatish bilan almashtirmang.',
 [('tolov(5000)','5000'),('tolov(5000,3)','15000'),('tolov(5000,0)','0')])
practice('practice-project','JavaScript konsol mini-loyihasi: o‘quvchilar ballari hisoboti',
 [('Talab','Dasturdan kutiladigan aniq natija: son, o‘rtacha va o‘tganlar miqdori.'),('Ma’lumot','Kirish uchun sonlardan iborat ballar massivi olinadi.'),('Hisoblash','Sikl orqali yig‘indi va shart orqali o‘tganlar soni topiladi.'),('Bo‘sh holat','Element bo‘lmaganda o‘rtacha hisoblanmaydi; nolga bo‘lish o‘rniga tushunarli xabar chiqariladi.'),('Qabul sinovi','Oldindan yozilgan kutilgan natija haqiqiy konsol natijasi bilan solishtiriladi.')],
 '41–67-mavzulardagi JavaScript asoslarini birlashtiring. Avval kirish va chiqishni yozing, keyin hisoblashni alohida funksiyalarga ajrating. Bo‘sh massivni tekshirgach sikl bilan yig‘indi va o‘tganlar miqdorini toping. Natija console.log orqali chiqadi. DOM, fetch, localStorage yoki veb-interfeys yaratish talab qilinmaydi; ular keyingi bo‘limlarda o‘rganiladi.',
 'const ballar = [40, 60, 80];\nlet summa = 0, otganlar = 0;\nfor (let i = 0; i < ballar.length; i++) {\n summa += ballar[i];\n if (ballar[i] >= 60) otganlar++;\n}\nif (ballar.length === 0) console.log("Ball yo‘q");\nelse { console.log(summa / ballar.length); console.log(otganlar); }','60\n2',
 'Konsolda ballar soni, o‘rtachasi, maksimumi va 60 balldan o‘tganlar miqdorini ko‘rsatadigan hisobot yarating. Hisoblashni funksiyalarga ajrating. Bo‘sh massiv, bitta ball va turli ballar bilan kamida uchta qabul sinovi yozing. HTML yoki DOM qo‘shmang.',
 'Tayyor misolni aynan ko‘chirish topshiriqni tugatmaydi: maksimum, alohida funksiyalar va mustaqil sinovlar ham bajarilishi kerak.',
 [('[40,60,80]: o‘rtacha / maksimum / o‘tganlar','60 / 80 / 2'),('[100]','100 / 100 / 1'),('[]','Ball yo‘q')])
def create_practice(n,key):
 a=P[key];esc=html.escape;criteria=['Talab va kutilgan natija oldindan yozilgan.','Faqat hozirgacha o‘tilgan JavaScript amallari qo‘llangan.','Kamida uchta kirish uchun kutilgan va haqiqiy natija solishtirilgan.','Xato yoki chegara holati izohlangan.','Kod tartibli va natijani qayta tekshirish mumkin.']
 sections=[('Dars maqsadi',a['task']),('Tayanch bilimlar',f'41–{n-1}-mavzularda tushuntirilgan JavaScript bilimlari.'),('Avval eslang',a['teaching'])]+a['terms']+[('Kod va amaliy namuna',a['code']),('Kutiladigan natija',a['expected']),('Mustaqil topshiriq',a['task']),('Tekshirish tartibi','Avval misolni bajaring. Keyin topshiriqni mustaqil yozing va quyidagi sinovlardan foydalaning.'),('Keng tarqalgan xato',a['pitfall'])]
 web=''.join('<h3>'+esc(h)+'</h3>'+('<pre class="lesson-code"><code>'+esc(v)+'</code></pre>' if h in ['Kod va amaliy namuna','Kutiladigan natija'] else '<p>'+esc(v)+'</p>') for h,v in sections)
 web+='<h3>Namunaviy qabul sinovlari</h3><table><tr><th>Kirish</th><th>Kutilgan natija</th></tr>'+''.join('<tr><td>'+esc(x)+'</td><td>'+esc(y)+'</td></tr>' for x,y in a['checks'])+'</table><h3>Topshirish mezonlari</h3>'+''.join('<p>'+esc(x)+'</p>' for x in criteria)
 web+='<h3>Qo‘shimcha o‘qish</h3><p><a href="https://developer.mozilla.org/en-US/docs/Web/JavaScript/Guide">MDN JavaScript qo‘llanmasi</a></p><h3>Atamalar</h3><table><tr><th>Atama</th><th>Izoh</th></tr>'+''.join('<tr><td>'+esc(x)+'</td><td>'+esc(y)+'</td></tr>' for x,y in a['terms'])+'</table>'
 slides=[{'title':esc(a['title']),'kind':'cover','body':['5-modul · JavaScript amaliyoti',f'{n}-mavzu · 2 soat']},{'title':'Tayanch bilimlar','body':[esc(a['teaching'])]}]+[{'title':esc(k),'body':[esc(v)]} for k,v in a['terms']]+[{'title':'Namuna kodi','body':[],'code':esc(a['code'])},{'title':'Kutiladigan natija','body':[esc(a['expected'])]},{'title':'Mustaqil topshiriq','body':[esc(a['task'])]},{'title':'Sinovlar','body':[esc(x+' → '+y) for x,y in a['checks']]},{'title':'Xatoni tekshiring','body':[esc(a['pitfall'])]},{'title':'Topshirish mezonlari','body':[esc(x) for x in criteria]}]
 qa=[[esc(k)+' nima?',esc(v)] for k,v in a['terms']]+[['Bu darsda qaysi oldingi bilimlar ishlatiladi?',esc(a['teaching'])],['Namuna natijasi qanday?',esc(a['expected'])],['Topshiriq qanday tekshiriladi?',esc('; '.join(x+' → '+y for x,y in a['checks']))],['Qaysi xatodan saqlanish kerak?',esc(a['pitfall'])],['Mustaqil topshiriqning talabi nima?',esc(a['task'])]]
 rng=random.Random(2800+n);quiz=[]
 for i,(k,v) in enumerate(a['terms']):
  opts=[v]+rng.sample([vv for kk,vv in a['terms'] if kk!=k],3);rng.shuffle(opts);quiz.append({'q':esc(k+' uchun to‘g‘ri izohni tanlang.'),'options':[esc(x) for x in opts],'correct':opts.index(v),'explanation':esc(k+': '+v)})
 for i,(x,y) in enumerate(a['checks']):
  opts=[y,'Kirish tekshirilmasdan natija tasdiqlanadi.','Dastur faqat HTML sahifa yaratishi kerak.','Kutilgan natijani yozish talab qilinmaydi.'];rng.shuffle(opts);quiz.append({'q':esc(x+' sinovida kutiladigan natija qaysi?'),'options':[esc(v) for v in opts],'correct':opts.index(y),'explanation':esc(x+' → '+y)})
 for question,right in [('Topshiriqdagi asosiy ehtiyot chorasi qaysi?',a['pitfall']),('Ushbu darsda ish qanday tasdiqlanadi?',criteria[2])]:
  opts=[right,'Faqat kod satrlari soni baholanadi.','Natija ko‘rilmasdan topshiriq tugatiladi.','Xato chiqsa barcha sinovlar bekor qilinadi.'];rng.shuffle(opts);quiz.append({'q':esc(question),'options':[esc(x) for x in opts],'correct':opts.index(right),'explanation':esc(right)})
 return {'number':n,'title':a['title'],'kind':'Amaliy','hours':2,'type':'amaliy','html':web,'slides':slides,'qa':qa,'quiz':quiz,'videos':[],'assignment':{'mode':'js','task':a['task'],'code':a['code'],'expected':a['expected'],'criteria':criteria,'teaching':a['teaching'],'prerequisite':f'41–{n-1}-mavzular.','stage':'JavaScript asoslari — mustahkamlash','runner':{'functionName':'__frontendExample','tests':[]}}}
def strip_glossary(markup):
 soup=BeautifulSoup(markup,'html.parser');terms=[]
 for h in list(soup.find_all(['h2','h3'])):
  if h.get_text(strip=True)=='Atamalar':
   table=h.find_next('table')
   if table:
    for row in table.find_all('tr')[1:]:
     cells=row.find_all(['td','th']);terms.append(tuple(c.get_text(' ',strip=True) for c in cells))
    wrapper=table.parent if table.parent.name=='div' and 'table-scroll' in table.parent.get('class',[]) else table;wrapper.decompose()
   h.decompose()
 return str(soup),terms
NEW={};NEW_CONTENT={}
for n,vs in units.items():
 if isinstance(vs[0],str):
  d=create_practice(n,vs[0]);a=P[vs[0]];NEW_CONTENT[n]={'terms':a['terms'],'mode':'js','code':a['code'],'task':a['task'],'pitfall':a['pitfall'],'expected':a['expected'],'teaching':a['teaching']}
 else:
  d=remap(deepcopy(old[vs[0]]));d['number']=n;NEW_CONTENT[n]=remap(deepcopy(content[vs[0]]))
  if 5<=n<=12:
   d['title']=html_titles[n-5]
   if len(vs)>1:
    chunks=[];glossary=[]
    for i,v in enumerate(vs):
     markup,terms=strip_glossary(remap(old[v]['html']));chunks.append('<section><h2>'+str(i+1)+'. '+html.escape(old[v]['title'])+'</h2>'+markup+'</section>');glossary+=terms
    glossary=list(dict.fromkeys(glossary));d['html']='<p class="notice">Ushbu 2 soatlik dars ikki bog‘liq mavzuni birlashtiradi. Avval ikkala tushuntirishni o‘qing, keyin quyidagi umumiy topshiriqni bajaring.</p>'+''.join(chunks)+'<h3>Atamalar</h3><table><tr><th>Atama</th><th>Izoh</th></tr>'+''.join('<tr><td>'+html.escape(k)+'</td><td>'+html.escape(v)+'</td></tr>' for k,v in glossary)+'</table>'
    d['slides']=[{'title':html.escape(d['title']),'body':['5-modul · HTML asoslari',f'{n}-mavzu · 2 soat'],'kind':'cover'}]
    for v in vs:d['slides']+=remap(deepcopy(old[v]['slides'][1:]))
    d['qa']=remap(old[vs[0]]['qa'][:5]+old[vs[1]]['qa'][:5]);d['quiz']=remap(old[vs[0]]['quiz'][:4]+[old[vs[0]]['quiz'][5]]+old[vs[1]]['quiz'][:4]+[old[vs[1]]['quiz'][5]])
    d['assignment']['teaching']=' '.join(remap(old[v]['assignment'].get('teaching','')) for v in vs)
    d['assignment']['task']=' '.join(remap(old[v]['assignment']['task']) for v in vs)
    d['assignment']['code']='\n'.join(old[v]['assignment']['code'] for v in vs)
    if n==5:d['assignment']['code']=old[5]['assignment']['code'].replace('</body>',old[6]['assignment']['code']+'</body>')
    if n==9:d['assignment']['code']='<header><h1>Texnikum blogi</h1><nav><a href="#xabar">Xabarlar</a> <a href="#aloqa">Aloqa</a></nav></header>\n<main><section id="xabar"><h2>Xabarlar</h2><article><h3>Ochiq dars</h3><p>Frontend taqdimoti.</p></article></section><section id="aloqa"><h2>Aloqa</h2><p>Namangan shahri</p></section></main>\n<footer>2026 — O‘quv loyiha</footer>'
    if n==10:d['assignment']['code']='<form><label for="ism">Ism</label><input id="ism" name="ism" required>\n<label for="email">Email</label><input id="email" name="email" type="email" required>\n<label for="kurs">Kurs</label><select id="kurs" name="kurs"><option>HTML</option><option>Keyingi kurs</option></select>\n<label><input type="radio" name="vaqt" value="ertalab"> Ertalab</label><label><input type="radio" name="vaqt" value="kechqurun"> Kechqurun</label>\n<label><input type="checkbox" name="rozilik"> Shartlarni o‘qidim</label><label for="izoh">Izoh</label><textarea id="izoh" name="izoh"></textarea><button type="submit">Yuborish</button></form>'
    NEW_CONTENT[n].update(terms=glossary,code=d['assignment']['code'],task=d['assignment']['task'],teaching=d['assignment']['teaching'],expected=d['assignment']['expected'])
   d['kind']='Nazariy-amaliy' if n in [5,12] else 'Amaliy'
   d['assignment']['prerequisite']='1–4-mavzular: boshlang‘ich tushunchalar va muharrir.' if n==5 else f'5–{n-1}-mavzularda o‘tilgan HTML bilimlari.'
   d['assignment']['stage']='HTML asoslari — faqat HTML'
  # Every reused cover must display its new lesson number.
  if d.get('slides'):d['slides'][0].update(title=html.escape(d['title']),body=[f'5-modul · {n}-mavzu',d['kind']+' · 2 soat'])
 d['number']=n;d['hours']=2;d['download']=f'materials/m5-r28-{n:03d}-konspekt.docx';d['sourceTopicNumbers']=[v for v in vs if isinstance(v,int)]
 if d.get('assignment'):
  d['assignment']['draftId']='r28-'+str(n)
  d['assignment']['bootstrap']=any(v in [43,44,45,46,99] for v in vs)
 if n==8:
  def update_hours(v):
   if isinstance(v,str):return v.replace('<td>HTML</td><td>28</td>','<td>HTML</td><td>16</td>').replace('&lt;td&gt;HTML&lt;/td&gt;&lt;td&gt;28&lt;/td&gt;','&lt;td&gt;HTML&lt;/td&gt;&lt;td&gt;16&lt;/td&gt;')
   if isinstance(v,list):return [update_hours(x) for x in v]
   if isinstance(v,dict):return {k:update_hours(x) for k,x in v.items()}
   return v
  d=update_hours(d);NEW_CONTENT[n]=update_hours(NEW_CONTENT[n])
 NEW[n]=d
# Put the combined task into the reading material as well as the workspace.
for n in range(5,11):
 d=NEW[n];soup=BeautifulSoup(d['html'],'html.parser');gloss=soup.find(['h2','h3'],string='Atamalar')
 extra=BeautifulSoup('<h3>Umumiy mustaqil topshiriq</h3><p>'+html.escape(d['assignment']['task'])+'</p><h3>Umumiy namuna</h3><pre class="lesson-code"><code>'+html.escape(d['assignment']['code'])+'</code></pre>','html.parser')
 for item in list(extra.contents):gloss.insert_before(item)
 d['html']=str(soup)
groups=[]
for g in BASE['plan']['groups']:
 number=g['number'];ns=range(1,5) if number==1 else range(5,13) if number==2 else range(13,31) if number==3 else range(31,41) if number==4 else range(41,69) if number==5 else [number_map[t['number']] for t in g['topics']]
 groups.append({'number':number,'name':g['name'],'hours':len(list(ns))*2,'topics':[{k:NEW[n][k] for k in ['number','title','kind','hours']} for n in ns]})
plan={'name':BASE['plan']['name'],'hours':220,'paras':110,'revision':'r28','groups':groups}
(ROOT/'authoring/frontend-plan.json').write_text(json.dumps(plan,ensure_ascii=False,indent=2));(ROOT/'authoring/frontend-content.json').write_text(json.dumps(NEW_CONTENT,ensure_ascii=False,indent=2));(ROOT/'authoring/frontend-number-map.json').write_text(json.dumps(number_map,indent=2))
(DIST/'frontend-data.js').write_text('const frontendName='+json.dumps(plan['name'],ensure_ascii=False)+';\nconst frontendGroups='+json.dumps(groups,ensure_ascii=False)+';\nconst frontendLessons=frontendGroups.flatMap(g=>g.topics);\nconst frontendData=Object.fromEntries(frontendLessons.map(t=>[t.number,t]));\n')
video_selection=ROOT/'authoring/frontend-video-selection.json'
if video_selection.exists():
 for number,videos in json.loads(video_selection.read_text()).items():NEW[int(number)]['videos']=videos
(DIST/'frontend-resources.js').write_text('const frontendResources='+json.dumps(NEW,ensure_ascii=False)+';\nObject.assign(frontendData,frontendResources);\n')

def styled_document():
 d=Document();sec=d.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11);sec.top_margin=sec.bottom_margin=Inches(.7);sec.left_margin=sec.right_margin=Inches(.75)
 for st in d.styles:
  if st.type in [1,2]:st.font.name='Calibri';st.font.size=Pt(14);st.font.color.rgb=RGBColor(0,0,0)
  if st.type==1:
   st.paragraph_format.line_spacing=1;st.paragraph_format.space_after=Pt(5)
   if st.element.pPr is not None:
    for b in list(st.element.pPr.findall(qn('w:pBdr'))):st.element.pPr.remove(b)
 for nm in ['Title','Heading 1','Heading 2']:
  st=d.styles[nm];st.font.bold=True;st.paragraph_format.keep_with_next=True;st.paragraph_format.space_before=Pt(8)
 d.core_properties.author='Ubaydullayev Abdukarim';return d

def format_table(tb):
 if len(tb.columns)==2:
  tb.autofit=False
  for col,w in zip(tb.columns,[1.8,5.2]):col.width=Inches(w)
  for row in tb.rows:
   for cell,w in zip(row.cells,[1.8,5.2]):cell.width=Inches(w)
 for i,row in enumerate(tb.rows):
  pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
  if i==0:pr.append(OxmlElement('w:tblHeader'))
  for cell in row.cells:
   cp=cell._tc.get_or_add_tcPr();border=OxmlElement('w:tcBorders')
   for side in ['top','left','bottom','right']:
    e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');border.append(e)
   cp.append(border)
   if i==0:e=OxmlElement('w:shd');e.set(qn('w:fill'),'DCE6F1');cp.append(e)
   for par in cell.paragraphs:
    par.paragraph_format.line_spacing=1
    if len(tb.columns)==2:par.paragraph_format.keep_with_next=i<len(tb.rows)-1
    for r in par.runs:r.font.name='Calibri';r.font.size=Pt(14);r.bold=i==0

def finalize(d,path):
 for par in d.paragraphs:
  par.paragraph_format.line_spacing=1
  for r in par.runs:r.font.name='Calibri';r.font.size=Pt(14)
 d.save(path)

for n,data in NEW.items():
 if '--only-plan' in sys.argv:break
 d=styled_document();d.core_properties.title=data['title'];g=next(g for g in groups if any(t['number']==n for t in g['topics']))
 d.add_paragraph('Web interfeysini ishlab chiqish Frontend',style='Title');d.add_paragraph(f'5 modul  {g["number"]} bo‘lim  {n} mavzu  {data["kind"]}  2 soat');d.add_paragraph(data['title'],style='Title');d.add_paragraph('Namangan shahar 3-sonli texnikum\nTayyorlovchi Ubaydullayev Abdukarim\n2026–2027 o‘quv yili')
 soup=BeautifulSoup(data['html'],'html.parser')
 for el in soup.find_all(['h2','h3','h4','p','pre','table']):
  if el.find_parent(['pre','table']):continue
  if el.name=='table':
   rows=el.find_all('tr');cols=max(len(r.find_all(['td','th'])) for r in rows);tb=d.add_table(rows=0,cols=cols)
   for row in rows:
    cells=tb.add_row().cells
    for i,cell in enumerate(row.find_all(['td','th'])):cells[i].text=cell.get_text(' ',strip=True)
   format_table(tb)
  elif el.name in ['h2','h3','h4']:d.add_heading(el.get_text(' ',strip=True),1)
  else:
   txt=el.get_text('\n' if el.name=='pre' else ' ',strip=True)
   if el.name=='p':
    for link in el.find_all('a'):
     if link.get('href','').startswith('https:'):txt+='\n'+link['href']
   par=d.add_paragraph(txt)
   if el.name=='pre':par.paragraph_format.keep_together=True
 finalize(d,DIST/data['download'])
# Updated teaching allocation; keep the uploaded source plan separately.
d=styled_document();d.core_properties.title='Frontend — yangilangan o‘quv taqsimoti';d.add_paragraph('Frontend — yangilangan o‘quv taqsimoti',style='Title');d.add_paragraph('Namangan shahar 3-sonli texnikum\nTayyorlovchi Ubaydullayev Abdukarim\n2026–2027 o‘quv yili');d.add_paragraph('Jami 220 soat — 110 para. Har bir para 2 soat. HTML: 16 soat — 8 para. JavaScript asoslari: 56 soat — 28 para. Qolgan bo‘limlarning soatlari saqlangan.');d.add_paragraph('HTMLdagi yaqin mavzular birlashtirildi. JavaScriptga o‘tilgan bilimlarni mustahkamlovchi oltita qo‘shimcha amaliy mashg‘ulot kiritildi.')
for g in groups:
 d.add_heading(f'{g["number"]}. {g["name"]} — {g["hours"]} soat, {len(g["topics"])} para',1);tb=d.add_table(rows=1,cols=3)
 tb.autofit=False
 for col,w in zip(tb.columns,[.5,5.8,.7]):col.width=Inches(w)
 for c,v in zip(tb.rows[0].cells,['№','Mavzu','Soat']):c.text=v
 for t in g['topics']:
  for c,v in zip(tb.add_row().cells,[str(t['number']),t['title'],'2']):c.text=v
 for row in tb.rows:
  for cell,w in zip(row.cells,[.5,5.8,.7]):cell.width=Inches(w)
 format_table(tb)
finalize(d,DIST/'materials/m5-frontend-yangilangan-reja.docx')
assert len(NEW)==110 and [g['hours'] for g in groups]==[8,16,36,20,56,20,8,16,40]
print('Built 110 lessons; HTML 16h/8 paras; JavaScript 56h/28 paras; total 220h. Regenerated 110 DOCX and revised plan.')
